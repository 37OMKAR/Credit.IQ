"""Read any document an officer uploads: PDF (digital / scanned), phone photos (JPG/PNG),
Word, Excel/CSV and text. Returns page-wise text so every module can use it."""
import io

from . import pdf_reader as PR

try:
    import pymupdf as fitz
except ImportError:  # older PyMuPDF
    import fitz

IMAGE_EXT = (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp")


def image_to_pdf(data: bytes, name: str) -> bytes:
    """A photographed balance sheet becomes a one-page PDF so the same OCR path handles it."""
    ext = name.lower().rsplit(".", 1)[-1]
    img = fitz.open(stream=data, filetype=ext)
    pdf_bytes = img.convert_to_pdf()
    img.close()
    return pdf_bytes


def images_to_pdf(items):
    """Several photos (pages of one statement) -> one PDF, in upload order."""
    out = fitz.open()
    for name, data in items:
        out.insert_pdf(fitz.open(stream=image_to_pdf(data, name), filetype="pdf"))
    return out.tobytes()


def as_pdf_bytes(name, data):
    return image_to_pdf(data, name) if name.lower().endswith(IMAGE_EXT) else data


def read_any(name, data, ocr_mode="auto", max_pages=450, progress=None):
    low = name.lower()
    if low.endswith(".pdf") or low.endswith(IMAGE_EXT):
        pages = PR.read_pdf(as_pdf_bytes(name, data), ocr_mode=ocr_mode, max_pages=max_pages, progress=progress)
        return {"name": name, "pages": pages, "text": "\n".join(p["text"] for p in pages)}
    if low.endswith(".docx"):
        from docx import Document
        d = Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append("   ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
    elif low.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"--- sheet: {ws.title}")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i > 400:
                    break
                cells = ["" if v is None else str(v) for v in row]
                if any(cells):
                    parts.append("   ".join(cells))
        text = "\n".join(parts)
    else:  # csv / txt / anything text-like
        text = data.decode("utf-8", errors="ignore")
    return {"name": name, "pages": [{"page": 1, "text": text, "ocr": False}], "text": text}
