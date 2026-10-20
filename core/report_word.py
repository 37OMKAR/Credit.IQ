"""Detailed credit-intelligence report in Word (.docx)."""
from .forensics import score_text
import datetime as dt
import io

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import credit_tools as CT
from . import explain as EXPL
from . import ratios as R
from . import schema
from docx.enum.text import WD_BREAK

# house style: warm near-black ink, orange accent, cream panels (matches the app)
INK = RGBColor(0x1C, 0x19, 0x17)
NAVY = INK  # kept for older callers
ORANGE = RGBColor(0xF2, 0x65, 0x22)
ORANGE_TXT = RGBColor(0xC2, 0x41, 0x0C)   # darker orange for text on white (readable)
INK_HEX, ORANGE_HEX, CREAM_HEX, RULE_HEX = "1C1917", "F26522", "F7F3EE", "E5DCD1"
SERIES = ["#1C1917", "#F26522", "#B8A897", "#2a78d6"]
STATUS = {"good": "DCFCE7", "watch": "FEF3C7", "bad": "FEE2E2"}
SEV_FILL = {"High": "FEE2E2", "Medium": "FEF3C7", "Low": "ECFDF3"}


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _borders(t, color="E5DCD1", inside=True):
    """Hairline warm-grey borders instead of Word's black grid."""
    tblPr = t._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        if edge in ("left", "right", "insideV") or (edge.startswith("inside") and not inside):
            e.set(qn("w:val"), "nil")
        else:
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)
    mar = OxmlElement("w:tblCellMar")
    for side, w in (("top", 45), ("bottom", 45), ("left", 90), ("right", 90)):
        m = OxmlElement(f"w:{side}"); m.set(qn("w:w"), str(w)); m.set(qn("w:type"), "dxa"); mar.append(m)
    tblPr.append(mar)


def _table(doc, header, rows, widths=None, font_size=8.5, header_fill="1C1917"):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    _borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(str(h))
        run.bold, run.font.size = True, Pt(font_size)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(c, header_fill)
    for row in rows:
        fills = None
        if isinstance(row, tuple) and len(row) == 2 and isinstance(row[1], dict):
            row, fills = row
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run("" if v is None else str(v))
            run.font.size = Pt(font_size)
            if i > 0 and isinstance(v, str) and (v[:1].isdigit() or v[:1] in "(-" or v == "n.a."):
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if fills and i in fills and fills[i]:
                _shade(cells[i], fills[i])
    if widths:
        t.autofit = False
        tblPr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblPr.append(lay)
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


def _num(v):
    if v is None:
        return "-"
    return f"({abs(v):,.2f})" if v < 0 else f"{v:,.2f}"


def _chart_financials(bundle):
    years = list(reversed(bundle["years"]))
    d = bundle["data"]
    keys = [("revenue", "Revenue"), ("ebitda", "EBITDA"), ("pat", "PAT")]
    fig, ax = plt.subplots(figsize=(7, 3.0), dpi=160)
    w = 0.8 / len(keys)
    for i, (k, lbl) in enumerate(keys):
        vals = [schema.g(d[y], k) for y in years]
        xs = [j + (i - (len(keys) - 1) / 2) * w for j in range(len(years))]
        ax.bar(xs, vals, width=w * 0.92, color=SERIES[i], label=lbl, zorder=3)
    ax.set_xticks(range(len(years)), years)
    ax.set_ylabel("INR Crore", fontsize=8)
    ax.set_title("Revenue, EBITDA and PAT", fontsize=10, loc="left", color="#0b0b0b")
    _style(ax)
    ax.legend(fontsize=7, frameon=False, ncol=3, loc="upper left")
    return _png(fig)


def _chart_leverage(bundle):
    years = list(reversed(bundle["years"]))
    r = bundle["ratios"]
    fig, ax = plt.subplots(figsize=(7, 2.8), dpi=160)
    for i, (k, lbl) in enumerate([("de", "Debt / Equity"), ("tol_tnw", "TOL / TNW"), ("debt_ebitda", "Debt / EBITDA")]):
        vals = [r[y].get(k) for y in years]
        ax.plot(years, [v if v is not None else float("nan") for v in vals], marker="o", markersize=5, linewidth=2, color=SERIES[i], label=lbl, zorder=3)
        if vals and vals[-1] is not None:
            ax.annotate(f"{vals[-1]:.2f}x", (len(years) - 1, vals[-1]), textcoords="offset points", xytext=(6, 0), fontsize=7, color="#52514e")
    ax.set_title("Leverage (times)", fontsize=10, loc="left", color="#0b0b0b")
    _style(ax)
    ax.legend(fontsize=7, frameon=False, ncol=3, loc="upper left")
    return _png(fig)


def _chart_flags(flags):
    cats = {}
    for f in flags:
        cats.setdefault(f["category"], {"High": 0, "Medium": 0, "Low": 0})[f["severity"] if f["severity"] in ("High", "Medium", "Low") else "Low"] += 1
    if not cats:
        return None
    names = sorted(cats, key=lambda c: -(cats[c]["High"] * 3 + cats[c]["Medium"] * 2 + cats[c]["Low"]))[:10]
    fig, ax = plt.subplots(figsize=(7, 0.35 * len(names) + 1.0), dpi=160)
    left = [0] * len(names)
    for sev, col in [("High", "#d03b3b"), ("Medium", "#fab219"), ("Low", "#0ca30c")]:
        vals = [cats[n][sev] for n in names]
        ax.barh(names, vals, left=left, color=col, label=sev, height=0.6, edgecolor="white", linewidth=1.5, zorder=3)
        left = [a + b for a, b in zip(left, vals)]
    ax.invert_yaxis()
    ax.set_title("Red flags by category and severity", fontsize=10, loc="left", color="#0b0b0b")
    _style(ax, grid_axis="x")
    ax.legend(fontsize=7, frameon=False, ncol=3, loc="lower right")
    return _png(fig)


def _style(ax, grid_axis="y"):
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.spines["bottom"].set_color("#c3c2b7")
    ax.tick_params(labelsize=7, colors="#52514e", length=0)
    ax.grid(axis=grid_axis, color="#e6e5e0", linewidth=0.6, zorder=0)


def _png(fig):
    buf = io.BytesIO()
    fig.tight_layout()
    fig.savefig(buf, format="png")
    plt.close(fig)
    buf.seek(0)
    return buf


def _para(doc, text, size=10, italic=False, bold=False, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text or "")
    run.font.size, run.italic, run.bold = Pt(size), italic, bold
    if color:
        run.font.color.rgb = color
    return p


def _bullets(doc, items, size=10):
    for it in items or []:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(str(it))
        run.font.size = Pt(size)




# ====================================================================== professional memo layout
from . import analyst as AN

FONT_BODY = "Calibri"
GREY = RGBColor(0x59, 0x59, 0x59)


def _field(paragraph, instr):
    """Insert a Word field (TOC, PAGE, NUMPAGES)."""
    run = paragraph.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = instr
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "Right-click and choose 'Update field' to build the table of contents." if "TOC" in instr else "1"
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    run._r.append(b); run._r.append(i); run._r.append(s); run._r.append(t); run._r.append(e)


def _p(doc, text, size=10.5, bold=False, italic=False, color=None, space_after=6, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text or "")
    r.font.size, r.bold, r.italic = Pt(size), bold, italic
    if color:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    return p


def _paras(doc, items, size=10.5):
    for t in items or []:
        if t:
            p = _p(doc, t, size)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _blt(doc, items, size=10):
    for it in items or []:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(str(it))
        r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(2)


def _note(doc, text):
    _p(doc, text, 8.5, italic=True, color=GREY, space_after=8)


def _fmtv(v, kind="cr"):
    if v is None:
        return "-"
    if kind == "pct":
        return "n.m." if abs(v) > 5 else f"{v * 100:,.1f}%"
    return f"({abs(v):,.2f})" if v < 0 else f"{v:,.2f}"


def style_headings(doc, sizes=(16, 12.5, 11)):
    """Ink headings; Heading 1 gets an orange rule underneath."""
    for (name, sz) in zip(("Heading 1", "Heading 2", "Heading 3"), sizes):
        stl = doc.styles[name]
        stl.font.name = FONT_BODY
        rf = stl.element.rPr.find(qn("w:rFonts"))
        if rf is not None:
            for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rf.attrib.pop(qn(att), None)
            rf.set(qn("w:ascii"), FONT_BODY); rf.set(qn("w:hAnsi"), FONT_BODY)
        stl.font.size = Pt(sz)
        stl.font.bold = True
        stl.font.color.rgb = INK
        stl.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        stl.paragraph_format.space_after = Pt(6 if name == "Heading 1" else 4)
    pPr = doc.styles["Heading 1"].element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "12"); bot.set(qn("w:space"), "4"); bot.set(qn("w:color"), ORANGE_HEX)
    bdr.append(bot); pPr.append(bdr)


def _cell_text(cell, parts, align=None):
    """parts: list of (text, size, bold, RGBColor)."""
    cell.text = ""
    first = True
    for text, size, bold, color in parts:
        para = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        r = para.add_run(text)
        r.font.size, r.bold = Pt(size), bold
        r.font.color.rgb = color
        para.paragraph_format.space_after = Pt(2)
        if align:
            para.alignment = align


def cover_band(doc, kicker, title, subtitle):
    """Dark full-width band with orange kicker, large title and grey subtitle."""
    t = doc.add_table(rows=1, cols=1)
    _borders(t, color=INK_HEX, inside=False)
    c = t.rows[0].cells[0]
    c.width = Cm(17)
    _shade(c, INK_HEX)
    _cell_text(c, [(" ", 6, False, INK), (kicker, 10, True, ORANGE), (title, 26, True, RGBColor(0xFF, 0xFF, 0xFF)),
                   (subtitle, 9.5, False, RGBColor(0xD6, 0xCC, 0xBF)), (" ", 6, False, INK)])
    _p(doc, "", space_after=6)


def cover_tiles(doc, tiles):
    """Row of KPI tiles: (label, value, sub, highlighted)."""
    t = doc.add_table(rows=1, cols=len(tiles))
    _borders(t, color="FFFFFF")
    for i, (lab, val, sub, hi) in enumerate(tiles):
        c = t.rows[0].cells[i]
        c.width = Cm(17 / len(tiles))
        _shade(c, ORANGE_HEX if hi else CREAM_HEX)
        _cell_text(c, [(lab, 8, True, INK if hi else RGBColor(0x6B, 0x62, 0x59)), (str(val), 15, True, INK),
                       (sub or " ", 8.5, False, INK if hi else RGBColor(0x6B, 0x62, 0x59))])
    _p(doc, "", space_after=6)


def _keyvalue(doc, rows, widths=(6.5, 10.5)):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    _borders(t)
    for a, b in rows:
        cells = t.add_row().cells
        for i, v in enumerate((a, b)):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(9.5)
            run.bold = i == 0
        _shade(cells[0], CREAM_HEX)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    return t


def build(bundle, narrative, officer="", branch=""):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT_BODY
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    st.font.size = Pt(10.5)
    style_headings(doc, (16, 12.5, 11))
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)

    years = bundle["years"]
    y0 = years[0]
    cols = list(reversed(years))
    d, r, fz = bundle["data"], bundle["ratios"], bundle["forensics"]
    company = bundle.get("company") or "Company"
    prop = bundle.get("proposal") or {}
    meta = bundle.get("meta") or {}
    notes = bundle.get("notes") or {}

    # header / footer with page numbers
    hp = sec.header.paragraphs[0]
    hp.text = f"STRICTLY CONFIDENTIAL  |  Credit Appraisal Memorandum  |  {company}"
    hp.runs[0].font.size, hp.runs[0].font.color.rgb = Pt(8), GREY
    fp = sec.footer.paragraphs[0]
    rr = fp.add_run("Prepared with Credit Intelligence Workbench - decision support; verify against audited statements.   Page ")
    rr.font.size, rr.font.color.rgb = Pt(7.5), GREY
    _field(fp, "PAGE")

    # ------------------------------------------------------------ cover
    nh = sum(1 for f in fz["flags"] if f["severity"] == "High")
    nm = sum(1 for f in fz["flags"] if f["severity"] == "Medium")
    cover_band(doc, "CREDIT APPRAISAL MEMORANDUM", company,
               f"{(bundle.get('statement_type') or 'standalone').title()} financial statements  ·  {cols[0]} to {cols[-1]}  ·  Amounts in ₹ crore")
    blocked_ = (bundle.get("integrity") or {}).get("status") == "BLOCKED"
    cover_tiles(doc, [("INTERNAL RATING", fz["grade"], ("PROVISIONAL - figures failed checks; " if blocked_ else "") + "score " + score_text(fz), False),
                      ("RECOMMENDATION", narrative.get("recommendation", "-"), "", False),
                      ("RED FLAGS", f"{nh} high", f"{nm} medium · {len(fz['flags']) - nh - nm} low", True)])
    cover = [
        ["Borrower", company],
        ["CIN", bundle.get("cin") or "-"],
        ["Proposal", f"{prop.get('facility')} - ₹{prop.get('amount', 0):,.2f} Cr" if prop.get("facility") else "Not specified"],
        ["Prepared by", f"{officer or '-'}" + (f", {branch}" if branch else "")],
        ["Date", dt.date.today().strftime("%d %B %Y")],
    ]
    _keyvalue(doc, cover, widths=(5, 12))
    chk_ = bundle.get("checks") or []
    if chk_:
        nf_ = sum(1 for c_ in chk_ if c_["result"] == "FAIL")
        nw_ = sum(1 for c_ in chk_ if c_["result"] == "WARN")
        _p(doc, "", space_after=4)
        ig_ = bundle.get("integrity") or {}
        if ig_:
            _p(doc, f"Case status: {ig_['status'].replace('_', ' ')}" + ("" if ig_["final_recommendation_allowed"] else
               " - no credit recommendation may be acted on") + f". {ig_.get('detail', '')}", 10, bold=True,
               color=RGBColor(0xB9, 0x1C, 0x1C) if ig_["status"] == "BLOCKED" else RGBColor(0xB4, 0x53, 0x09) if ig_["status"] != "READY" else RGBColor(0x15, 0x80, 0x3D),
               space_after=3)
        _p(doc, (f"DRAFT - {nf_} integrity check(s) FAILED. Correct the figures before relying on this memorandum. " if nf_ else "")
           + f"Logic and source checks: {len(chk_) - nf_ - nw_} passed, {nw_} warning(s), {nf_} failed - see Annexure D.",
           9.5, bold=bool(nf_), color=RGBColor(0xB9, 0x1C, 0x1C) if nf_ else GREY, space_after=4)
    _p(doc, "", space_after=10)
    _note(doc, "This memorandum is generated from the documents listed in Annexure C. Figures are extracted automatically and "
               "validated through tie-out checks (Section 2); the credit officer must verify material figures against the audited "
               "financial statements before sanction.")
    doc.add_page_break()

    # ------------------------------------------------------------ TOC
    _p(doc, "Contents", 16, bold=True, color=INK)
    toc_anchor = doc.add_paragraph()  # filled with the real section list once the document is complete
    doc.add_page_break()

    # ------------------------------------------------------------ 1. executive summary
    doc.add_heading("1. Executive Summary", 1)
    if narrative.get("_source") == "AI" and narrative.get("executive_summary"):
        _paras(doc, [narrative["executive_summary"]])
    _paras(doc, [AN.rating_rationale(bundle, narrative)])
    from . import credit_tools as CT
    from . import report_sections as RS
    P = CT.proposal(bundle)
    _p(doc, "Proposal at a glance", 10.5, bold=True, space_after=2)
    RS.proposal_snapshot(doc, bundle, P)
    _p(doc, "", space_after=4)
    doc.add_heading("1.1 Key credit metrics", 2)
    km = []
    for k in EXPL.SUMMARY_KEYS:
        v = r[y0].get(k)
        if v is None:
            continue
        stt = R.status(k, v)
        meta_r = R.RATIO_META[k]
        bench = ("≥ " if meta_r[5] == "high" else "≤ " if meta_r[5] == "low" else "") + (R.fmt(k, meta_r[6]) if meta_r[6] is not None else "")
        prev_v = r[years[1]].get(k) if len(years) > 1 else None
        km.append(([meta_r[1], R.fmt(k, prev_v) if prev_v is not None else "-", R.fmt(k, v), bench,
                    {"good": "Satisfactory", "watch": "Watch", "bad": "Weak"}.get(stt, "-")], {4: STATUS.get(stt)}))
    _table(doc, ["Metric", years[1] if len(years) > 1 else "Prior", y0, "Benchmark", "Assessment"], km,
           widths=[5.5, 2.6, 2.6, 2.8, 3.5], font_size=9)
    s_, c_ = AN.strengths_concerns(bundle)
    doc.add_heading("1.2 Key strengths", 2)
    _blt(doc, (narrative.get("strengths") if narrative.get("_source") == "AI" else None) or s_ or ["No metric is comfortably within benchmark."])
    doc.add_heading("1.3 Key concerns", 2)
    _blt(doc, c_ or ["No material concerns identified by the rule set."])
    doc.add_heading("1.4 Recommendation", 2)
    _paras(doc, [f"{narrative.get('recommendation', '-')}. {narrative.get('recommendation_rationale', '')}"])
    comments = bundle.get("comments") or []
    if comments:
        doc.add_heading("1.5 Credit officer's observations", 2)
        for cm in comments:
            p = doc.add_paragraph(style="List Bullet")
            a = p.add_run(f"{cm.get('topic', 'General')} ({cm.get('when', '')}): ")
            a.bold, a.font.size = True, Pt(9.5)
            b_ = p.add_run(cm.get("text", ""))
            b_.font.size = Pt(9.5)

    # ------------------------------------------------------------ 2. borrower & data
    doc.add_heading("2. Borrower Profile and Data Validation", 1)
    doc.add_heading("2.1 Borrower profile", 2)
    if narrative.get("_source") == "AI" and narrative.get("business_profile"):
        _paras(doc, [narrative["business_profile"]])
    _keyvalue(doc, [
        ["Name", company], ["CIN", bundle.get("cin") or "-"],
        ["Statements analysed", f"{(bundle.get('statement_type') or 'standalone').title()}, {', '.join(cols)}"],
        ["Statutory auditor", notes.get("auditor_name") or "-"],
        ["Audit opinion", AN._cap(notes.get("auditor_opinion") or "-")],
        ["Revenue (latest)", AN.cr(d[y0].get("revenue"))], ["Net worth (latest)", AN.cr(d[y0].get("total_equity"))],
        ["Funded debt (latest)", AN.cr(d[y0].get("total_debt"))],
    ])
    doc.add_heading("2.2 Data extraction and validation", 2)
    _paras(doc, ["Financial statements were located within the source documents, read (with optical character recognition where pages "
                 "were scanned) and mapped to a standard Schedule III spread. The following integrity checks were applied."])
    vrows = AN.data_validation(bundle)
    _table(doc, ["Check", "Result", "Detail"], [([a, b_, c__], {1: {"Passed": "DCFCE7", "Check": "FEE2E2", "Verify": "FEF3C7"}.get(b_)}) for a, b_, c__ in vrows],
           widths=[4.5, 2, 10.5], font_size=8.5)
    dual_ = meta.get("dual") or {}
    if dual_.get("rows"):
        _p(doc, "Second, independent reading", 10.5, bold=True, space_after=2)
        _paras(doc, [f"The statement pages ({', '.join(str(p_) for p_ in dual_.get('pages') or [])}) were read a second time directly from the page "
                     f"images by a separate model ({dual_.get('model')}); lines where the two readings differed were re-read by a third, "
                     f"OCR-based reader. {dual_['summary']['text']}."
                     + (" Corrected on the strength of two agreeing readers: " + "; ".join(dual_["corrections"]) + "." if dual_.get("corrections") else "")])
        open_ = [r_ for r_ in dual_["rows"] if not r_["status"].startswith(("Confirmed", "Corrected"))]
        if open_:
            _table(doc, ["Line", "Year", "Reader 1", "Reader 2", "Tie-breaker", "Status"],
                   [([r_["label"], r_["year"], _fmtv(r_["r1"]), _fmtv(r_["r2"]), _fmtv(r_["r3"]), r_["status"]], {5: "FEE2E2"}) for r_ in open_[:30]],
                   widths=[4.6, 1.6, 2.2, 2.2, 2.2, 4.2], font_size=8)
            _note(doc, "Lines above were not confirmed by two readers - verify them against the statement page before relying on the analysis.")
    elif (meta.get("dual_enabled") is False):
        _note(doc, "A second, independent reading of the statements was not performed (switched off or no OpenRouter key); "
                   "figures rest on one reading and the officer's check of the spread.")

    # ------------------------------------------------------------ 3. financial performance
    doc.add_heading("3. Financial Performance Analysis", 1)
    doc.add_heading("3.1 Income statement", 2)
    _paras(doc, AN.income_commentary(bundle))
    if narrative.get("_source") == "AI" and narrative.get("financial_performance"):
        _paras(doc, [narrative["financial_performance"]])
    doc.add_picture(_chart_financials(bundle), width=Cm(16.5))
    _note(doc, "Figure 1 - Revenue, EBITDA and PAT (₹ crore).")
    yrs, rows = AN.horizontal_table(bundle, ["revenue", "other_income", "cogs", "employee_cost", "other_expenses", "ebitda",
                                             "depreciation", "finance_cost", "pbt", "tax_expense", "pat"])
    _table(doc, ["Horizontal analysis (₹ Cr)"] + yrs + ["YoY change"],
           [[rw[0]] + [_fmtv(v) for v in rw[1:-1]] + [_fmtv(rw[-1], "pct")] for rw in rows], font_size=8.5)
    _note(doc, "Table - Income statement with year-on-year change for the latest year.")
    yrs, rows = AN.vertical_pl(bundle)
    _table(doc, ["Vertical analysis (% of revenue)"] + yrs, [[rw[0]] + [_fmtv(v, "pct") for v in rw[1:]] for rw in rows], font_size=8.5)
    _note(doc, "Table - Common-size income statement; each line expressed as a percentage of revenue.")

    doc.add_heading("3.2 Balance sheet", 2)
    _paras(doc, AN.balance_commentary(bundle))
    yrs, rows = AN.vertical_bs(bundle)
    _table(doc, ["Balance sheet composition (% of total assets)"] + yrs, [[rw[0]] + [_fmtv(v, "pct") for v in rw[1:]] for rw in rows], font_size=8.5)
    bs_rows = []
    for k, label, sec_, _ in schema.BALANCE_SHEET:
        if any(d[y].get(k) not in (None, 0) for y in cols):
            bs_rows.append([label] + [_fmtv(d[y].get(k)) for y in cols])
    doc.add_heading("3.3 Cash flow and quality of earnings", 2)
    _paras(doc, AN.cashflow_commentary(bundle))
    yrs, rows = AN.horizontal_table(bundle, ["pat", "depreciation", "cfo", "capex", "fcf", "cfi", "cff"])
    if rows:
        _table(doc, ["Cash flow (₹ Cr)"] + yrs + ["YoY change"], [[rw[0]] + [_fmtv(v) for v in rw[1:-1]] + [_fmtv(rw[-1], "pct")] for rw in rows], font_size=8.5)
    su = AN.sources_uses(bundle)
    if su:
        doc.add_heading("3.4 Sources and uses of funds", 2)
        _paras(doc, AN.sources_uses_commentary(su))
        n = max(len(su["sources"]), len(su["uses"]))
        rows = []
        for i in range(min(n, 12)):
            a = su["sources"][i] if i < len(su["sources"]) else ("", None, "")
            b2 = su["uses"][i] if i < len(su["uses"]) else ("", None, "")
            rows.append([f"{a[0]} ({a[2]})" if a[0] else "", _fmtv(a[1]) if a[1] is not None else "",
                         f"{b2[0]} ({b2[2]})" if b2[0] else "", _fmtv(b2[1]) if b2[1] is not None else ""])
        rows.append(["Total sources", _fmtv(su["total_sources"]), "Total uses", _fmtv(su["total_uses"])])
        _table(doc, [f"Sources of funds ({su['from']} to {su['to']})", "₹ Cr", "Uses of funds", "₹ Cr"], rows, widths=[6, 2.5, 6, 2.5], font_size=8.5)
        _note(doc, "Derived from movements between the two latest balance sheets. Increases in liabilities and decreases in assets are sources; the reverse are uses."
              + (f" Excluded because they were read for only one of the two years: {', '.join(su['excluded'])}." if su.get("excluded") else "")
              + (f" Sources and uses differ by ₹{abs(su['total_sources'] - su['total_uses']):,.2f} Cr - the statement does not fully reconcile; "
                 "check the spread before relying on it." if abs(su["total_sources"] - su["total_uses"]) > max(0.05, 0.005 * su["total_sources"]) else
                 " Sources and uses reconcile."))

    an = AN.anomalies(bundle)
    if an:
        doc.add_heading("3.5 Line-item anomalies", 2)
        _paras(doc, ["The following line items moved by more than 50% year on year and are material (above 2% of total assets or revenue). "
                     "Large unexplained movements are where errors, reclassifications and window-dressing are most often found; each should be "
                     "explained by management with supporting schedules."])
        _table(doc, ["Line item", years[1], y0, "Change", "Materiality"], an, widths=[5.5, 2.8, 2.8, 2.2, 3.7], font_size=8.5)

    # ------------------------------------------------------------ 4. ratio analysis
    doc.add_heading("4. Ratio Analysis", 1)
    _paras(doc, ["Ratios are computed from the standardised spread using a single formula set shared with the Excel workbook. "
                 "Benchmarks reflect common Indian bank credit norms and should be read alongside sector characteristics."])
    doc.add_picture(_chart_leverage(bundle), width=Cm(16.5))
    _note(doc, "Figure 2 - Leverage trend (times).")
    cat_text = {
        "Liquidity": lambda: f"Short-term liquidity: current ratio {AN.x(r[y0].get('current_ratio'))} and quick ratio {AN.x(r[y0].get('quick_ratio'))}; "
                             f"net working capital {AN.cr(r[y0].get('nwc_cr'))}.",
        "Leverage": lambda: f"Capital structure: debt-to-equity {AN.x(r[y0].get('de'))}, TOL/TNW {AN.x(r[y0].get('tol_tnw'))} and debt/EBITDA {AN.x(r[y0].get('debt_ebitda'))}. "
                            + ("TOL/TNW is above the 3x norm, driven by outside liabilities other than funded debt." if (r[y0].get('tol_tnw') or 0) > 3 and (r[y0].get('de') or 0) < 0.5 else ""),
        "Coverage": lambda: f"Debt service: interest coverage {AN.x(r[y0].get('icr'))}, EBITDA/interest {AN.x(r[y0].get('ebitda_int'))}, DSCR {AN.x(r[y0].get('dscr'))}"
                            + (" (DSCR requires the current maturities of long-term debt, which were not available)." if r[y0].get('dscr') is None else "."),
        "Profitability": lambda: f"Returns: EBITDA margin {AN.pct(r[y0].get('ebitda_margin'))}, PAT margin {AN.pct(r[y0].get('pat_margin'))}, ROE {AN.pct(r[y0].get('roe'))}, ROCE {AN.pct(r[y0].get('roce'))}.",
        "Efficiency": lambda: f"Asset utilisation: asset turnover {AN.x(r[y0].get('asset_turnover'))}; debtor days {AN._d(r[y0].get('debtor_days'))}, inventory days {AN._d(r[y0].get('inventory_days'))}, creditor days {AN._d(r[y0].get('creditor_days'))}.",
        "Cash Flow Quality": lambda: f"Cash conversion: CFO/EBITDA {AN.x(r[y0].get('cfo_ebitda'))}, CFO/PAT {AN.x(r[y0].get('cfo_pat'))}, free cash flow {AN.cr(r[y0].get('fcf_cr'))}.",
    }
    cats = list(dict.fromkeys(rt[2] for rt in R.RATIOS))
    for ci, cat in enumerate(cats, 1):
        buf = []
        for key, name, cat2, *_rest in R.RATIOS:
            if cat2 != cat:
                continue
            vals = [R.fmt(key, r[y].get(key)) for y in cols]
            if all(v == "n.a." for v in vals):
                continue
            meta_r = R.RATIO_META[key]
            stt = R.status(key, r[y0].get(key))
            bench = ("≥ " if meta_r[5] == "high" else "≤ " if meta_r[5] == "low" else "") + (R.fmt(key, meta_r[6]) if meta_r[6] is not None else "-")
            buf.append(([name] + vals + [bench, {"good": "Satisfactory", "watch": "Watch", "bad": "Weak"}.get(stt, "-")],
                        {len(cols) + 2: STATUS.get(stt)}))
        if not buf:
            continue
        doc.add_heading(f"4.{ci} {cat}", 2)
        if cat in cat_text:
            _paras(doc, [cat_text[cat]()])
        _table(doc, ["Ratio"] + cols + ["Benchmark", "Assessment"], buf, widths=[5.2] + [2.2] * len(cols) + [2.4, 2.6], font_size=8.5)

    sector = meta.get("sector")
    comp = AN.sector_comparison(bundle, sector) if sector else None
    if comp:
        doc.add_heading(f"4.{len(cats) + 1} Sector comparison - {sector}", 2)
        weaker = [r_[0] for r_ in comp if r_[3].startswith("Weaker")]
        _paras(doc, [f"Against the bank-approved norms for the {sector.lower()} sector, the company is weaker on "
                     + (", ".join(w.lower() for w in weaker) if weaker else "none of the measures compared") + "."])
        _table(doc, ["Measure", f"Company {y0}", "Sector norm", "Position"],
               [(r_, {3: "DCFCE7" if not r_[3].startswith("Weaker") else "FEE2E2"}) for r_ in comp], widths=[4.5, 3, 3, 6.5], font_size=9)
        _note(doc, "Sector norms are the bank-approved values maintained by the administrator (Admin › Bank parameters).")

    RS.dupont_section(doc, bundle, num=f"4.{len(cats) + (2 if comp else 1)}")

    # ------------------------------------------------------------ 5. working capital
    doc.add_heading("5. Working Capital Assessment and Bank Finance", 1)
    _wc_miss = CT.core_missing(bundle, ("revenue", "total_current_assets", "total_current_liabilities"))
    if _wc_miss:
        _paras(doc, [f"Not computed - the spread is missing: {', '.join(_wc_miss)}. The operating cycle, net working capital and maximum "
                     "permissible bank finance would otherwise be computed on blanks treated as zero."])
    else:
        wtxt, wtab = AN.working_capital(bundle)
        _paras(doc, wtxt)
        if narrative.get("_source") == "AI" and narrative.get("working_capital"):
            _paras(doc, [narrative["working_capital"]])
        _table(doc, ["Assessment of maximum permissible bank finance", "₹ Cr"], wtab, widths=[12, 5], font_size=9)
    RS.wc_section(doc, bundle, P, num="5")

    # ------------------------------------------------------------ 6. repayment capacity & stress
    doc.add_heading("6. Repayment Capacity, Projections and Stress Test", 1)
    nx = RS.repayment_section(doc, bundle, P, num="6")
    doc.add_heading(f"6.{nx} Stress test on latest audited year", 2)
    srows, worst = AN.stress_test(bundle)
    if CT.core_missing(bundle, ("revenue", "pbt", "pat")):
        _paras(doc, ["Not computed - revenue or profit lines are missing from the spread."])
    elif (d[y0].get("total_debt") or 0) < 0.01 and all(v_ == "n.a." for r_ in srows for v_ in r_[1:3]):
        _paras(doc, ["The company has no funded debt, so the interest-rate and debt-service stress tests do not apply. The relevant "
                     "stress is on liquidity - see the next-year liquidity check below."])
    else:
        _paras(doc, ["The latest year's earnings were stressed for a 20% fall in EBITDA and a 200 basis point rise in the cost of funded debt, "
                     "individually and together, holding other items constant (tax at 25.17%)."
                     + (f" Under the combined scenario interest coverage falls to {worst[0]:.2f}x." if worst[0] is not None else "")])
        _table(doc, ["Scenario", "Interest coverage", "DSCR", "Debt / EBITDA"], srows, widths=[5, 4, 4, 4], font_size=9)
    doc.add_heading(f"6.{nx + 1} Next-year liquidity check", 2)
    _lq_miss = CT.core_missing(bundle)
    if _lq_miss:
        _paras(doc, [f"Not computed - the spread is missing: {', '.join(_lq_miss)}."])
    else:
        prow, pnote = AN.projection(bundle)
        _paras(doc, [pnote + " Assumptions: revenue grows at the latest year's rate (capped at 25%), EBITDA margin and working-capital "
                     "days are held constant, tax at 25.17%."])
        _table(doc, ["Projection item", "₹ Cr / %"], prow, widths=[10, 6], font_size=9)

    # ------------------------------------------------------------ 7. forensic
    doc.add_heading("7. Forensic Accounting Review", 1)
    _paras(doc, ["The forensic review tests for earnings manipulation, diversion of funds and undisclosed liabilities. It combines "
                 "statistical models (Beneish M-score, Piotroski F-score, Altman Z''), fund-flow and end-use tests, working-capital "
                 "anomalies, cash-yield and cost-of-debt tests, and a review of the auditor's report, CARO annexure and notes to accounts."])
    if narrative.get("_source") == "AI" and narrative.get("forensic_assessment"):
        _paras(doc, [narrative["forensic_assessment"]])
    ch = _chart_flags(fz["flags"])
    if ch:
        doc.add_picture(ch, width=Cm(16.5))
        _note(doc, "Figure 3 - Observations by area and severity.")
    order = {"High": 0, "Medium": 1, "Low": 2}
    has_rm = any(f.get("remark") or f.get("status", "Open") != "Open" for f in fz["flags"])
    frows = []
    for i, f in enumerate(sorted(fz["flags"], key=lambda x_: order.get(x_["severity"], 3)), 1):
        row = [str(i), f["severity"], f["title"], f["evidence"][:350], f.get("why", ""), f["action"]]
        if has_rm:
            row.append(f"{f.get('status', 'Open')}" + (f": {f['remark']}" if f.get("remark") else ""))
        frows.append((row, {1: SEV_FILL.get(f["severity"])}))
    if frows:
        hdr = ["#", "Severity", "Observation", "Evidence", "Implication", "Action required"] + (["Officer disposition"] if has_rm else [])
        wid = [0.7, 1.5, 2.8, 4.2, 3.2, 3.0] + ([2.4] if has_rm else [])
        _table(doc, hdr, frows, widths=wid, font_size=7.5)
    else:
        _paras(doc, ["No observations were triggered."])
    na_ = fz.get("not_assessed") or []
    if na_:
        _p(doc, "Tests not assessed - the figures they need were not read", 10, bold=True, space_after=2)
        _paras(doc, ["These tests were not run, so the absence of a flag here is NOT a clean result. Enter the missing lines on Spread review "
                     "and regenerate."])
        _table(doc, ["Forensic test", "Inputs not read"], [[t["test"], ", ".join(t["missing"])] for t in na_], widths=[8, 9], font_size=8)
    doc.add_heading("7.1 Statistical models", 2)
    from .forensics import MODEL_NOTE
    models_ = fz.get("models") or {}
    if models_:
        _table(doc, ["Model", "Status", "Note"], [([m, v["status"], (v.get("note") or "") + (": " + ", ".join(v["missing"]) if v.get("missing") else "")],
                                                   {1: {"CALCULATED": "DCFCE7", "PARTIAL": "FEF3C7", "NOT COMPUTED": "FEE2E2"}.get(v["status"])})
                                                  for m, v in models_.items()], widths=[4, 3, 10], font_size=8.5)
        _note(doc, MODEL_NOTE)
    bm, pf, az = fz.get("beneish"), fz.get("piotroski"), fz.get("altman_z")
    _paras(doc, [
        (f"Beneish M-score is {bm:.2f}" + (" - above the -1.78 threshold, placing the company in the zone associated with earnings manipulation; "
                                          "the highest component indices should be investigated." if bm > -1.78 else " - below the -1.78 threshold, not indicative of manipulation.")) if bm is not None else "Beneish M-score: NOT COMPUTED (see the status table above).",
        (f"Piotroski F-score is {pf}/9" + (" (strong)." if pf >= 7 else " (weak)." if pf <= 3 else " (average).")) if pf is not None else "",
        (f"Altman Z'' (non-manufacturing form; zones 1.1 / 2.6) is {az:.2f}" + (" - safe zone." if az > 2.6 else " - distress zone." if az < 1.1 else " - grey zone.")) if az is not None else "",
    ])
    bv = fz.get("beneish_vars") or {}
    if bv:
        _table(doc, ["Beneish index", "Value", "Reading"], [[k, f"{v:.3f}", "Deterioration" if v > 1.05 and k not in ("SGAI", "LVGI", "TATA") else "-"] for k, v in bv.items()],
               widths=[4, 3, 6], font_size=8.5)

    RS.ews_section(doc, bundle, P, num="7.2")
    RS.screening_section(doc, bundle, num="7.3")

    doc.add_heading("8. Auditor's Report, Notes and Contingent Liabilities", 1)
    if notes.get("available"):
        _keyvalue(doc, [
            ["Auditor", notes.get("auditor_name") or "-"], ["Opinion", AN._cap(notes.get("auditor_opinion") or "-")],
            ["Going-concern uncertainty", "Not verified - auditor's report not supplied" if notes.get("going_concern_uncertainty") is None
             and str(notes.get("auditor_opinion", "")).startswith("not verified") else ("Yes" if notes.get("going_concern_uncertainty") else "No")],
            ["Contingent liabilities", AN.cr(notes.get("contingent_liabilities_total")) if isinstance(notes.get("contingent_liabilities_total"), (int, float)) else "-"],
            ["Loans / advances to related parties", AN.cr(notes.get("loans_advances_to_related_parties")) if isinstance(notes.get("loans_advances_to_related_parties"), (int, float)) else "-"],
            ["Guarantees for related parties", AN.cr(notes.get("guarantees_given_for_related_parties")) if isinstance(notes.get("guarantees_given_for_related_parties"), (int, float)) else "-"],
        ])
        for title, k in [("Key audit matters", "key_audit_matters"), ("Emphasis of matter", "emphasis_of_matter"),
                         ("CARO observations", "caro_adverse_remarks"), ("Defaults / delays", "defaults_or_delays"), ("Contingent items", "contingent_items")]:
            if notes.get(k):
                _p(doc, title, 10.5, bold=True, space_after=2)
                _blt(doc, notes[k], 9.5)
        rps = notes.get("related_parties") or []
        if rps:
            _p(doc, "Material related-party transactions", 10.5, bold=True, space_after=2)
            _table(doc, ["Party", "Relationship", "Transaction", "₹ Cr"],
                   [[x_.get("name"), x_.get("relationship"), x_.get("transaction"), _fmtv(x_.get("amount"))] for x_ in rps[:30]], font_size=8)
        if notes.get("method") == "keyword scan":
            _note(doc, "Notes were reviewed by keyword scan (no AI key configured); an AI review gives a fuller reading of related-party and contingent-liability notes.")
    else:
        _paras(doc, ["The auditor's report and notes could not be located in the documents provided."])

    # ------------------------------------------------------------ market intel
    doc.add_heading("9. Market Intelligence", 1)
    news = bundle.get("news") or []
    brief = bundle.get("brief") or {}
    if not news:
        _paras(doc, ["No live news was retrieved for this memorandum. Re-run with 'Pull live news' enabled to include press coverage, "
                     "regulatory actions, rating changes and adverse-media screening."])
    else:
        neg = sum(1 for n_ in news if n_.get("sentiment") == "Negative")
        hi = sum(1 for n_ in news if n_.get("severity") == "High")
        _paras(doc, [f"Adverse-media and market screening covered {len(news)} articles from the last twelve months: {neg} negative in tone and "
                     f"{hi} of high severity (fraud, default, investigation or rating action)."])
        if brief.get("overview"):
            _paras(doc, [brief["overview"]])
            if brief.get("overview_source"):
                _note(doc, f"Source of the paragraph above: {brief['overview_source']}. Treat statements without a listed source as leads to confirm.")
        for title, k in [("Positives", "key_positives"), ("Concerns", "key_concerns"), ("Regulatory and legal", "regulatory_legal"),
                         ("Management and governance", "management_governance"), ("Items needing verification", "items_needing_verification")]:
            if brief.get(k):
                _p(doc, title, 10.5, bold=True, space_after=2)
                _blt(doc, brief[k], 9.5)
        cr_ = brief.get("credit_rating") or {}
        if cr_.get("rating"):
            _p(doc, f"External credit rating: {cr_.get('agency') or ''} {cr_.get('rating')} {cr_.get('outlook') or ''}"
                    + (f" ({cr_.get('date')})" if cr_.get("date") else "") + (f" - {cr_.get('url')}" if cr_.get("url") else ""), 10, bold=True)
        if (brief.get("credit_rating_unverified") or {}).get("rating"):
            cu_ = brief["credit_rating_unverified"]
            _note(doc, f"A web search suggested a rating of {cu_.get('agency') or ''} {cu_.get('rating')} but it could not be confirmed on a source page "
                       f"({cu_.get('verification')}); it is not used. Obtain the rating letter / rationale from the borrower.")
        sev_rank = {"High": 0, "Medium": 1, "Low": 2}
        top = sorted(news, key=lambda n_: (sev_rank.get(n_.get("severity"), 3), n_.get("date", "")))[:25]
        _table(doc, ["Date", "Headline", "Source", "Tone", "Category"],
               [([n_.get("date"), n_.get("title"), n_.get("source"), n_.get("sentiment"), n_.get("category")],
                 {3: {"Negative": "FEE2E2", "Positive": "DCFCE7"}.get(n_.get("sentiment"))}) for n_ in top],
               widths=[2.0, 8.0, 2.8, 1.8, 2.6], font_size=7.5)
        _note(doc, "Headlines come from Google News or from web-search results whose source page was opened and checked; tone and category are AI tags.")
    if brief.get("unverified_web"):
        _note(doc, f"{len(brief['unverified_web'])} web-search item(s) were rejected because the source page could not be opened or did not carry the "
                   "story about this company - they are not used anywhere in this memorandum.")

    if brief.get("web_sources"):
        # keep only sources that are about this company (search engines also return look-alike names and generic pages)
        from . import watchlist as _WL
        import re as _re
        toks = _WL.distinctive(company)
        cin_ = (bundle.get("cin") or "").upper()

        def _about(c_):
            t_ = _WL.norm((c_.get("title") or "") + " " + _re.sub(r"[-_/.%]", " ", c_.get("url") or ""))
            words = set(t_.split())
            hit = sum(1 for k_ in toks if k_ in words)
            return (cin_ and cin_ in (c_.get("url") or "").upper()) or (toks and hit >= min(2, len(toks)))
        rel = [c_ for c_ in brief["web_sources"] if _about(c_)]
        if rel:
            _p(doc, "Web sources consulted", 10.5, bold=True, space_after=2)
            _blt(doc, [f"{c_.get('title', '')[:90]} - {c_.get('url')}" for c_ in rel[:12]], 8.5)
        if len(rel) < len(brief["web_sources"]):
            _note(doc, f"{len(brief['web_sources']) - len(rel)} other search result(s) were about similarly named or unrelated entities and are not relied on.")

    # ------------------------------------------------------------ ESG / PCAF
    doc.add_heading("10. ESG, BRSR and Financed Emissions", 1)
    b_ = bundle.get("brsr") or {}
    we = (meta.get("web_esg") or {})
    if we:
        _p(doc, "Web check of BRSR filing and disclosures", 10.5, bold=True, space_after=2)
        _paras(doc, [("A Business Responsibility and Sustainability Report was found for the company"
                      + (f" (years: {', '.join(we.get('brsr_years_filed') or [])})" if we.get("brsr_years_filed") else "") + "."
                      if we.get("brsr_filed") else "No BRSR filing was found on the web for the company.")
                     + f" BRSR applicability: {we.get('brsr_applicable') or 'unknown'}. Emissions "
                     + ("are declared" if we.get("emissions_declared") else "were not found as declared")
                     + f"; GHG assurance: {we.get('ghg_assurance') or 'not stated'}." + (f" {we.get('notes')}" if we.get("notes") else "")])
        rows_ = [[lab, (f"{v:,.2f}" if isinstance(v, (int, float)) else str(v))] for lab, v in
                 [("Listed / symbol", f"{'Yes' if we.get('listed') else 'No'} {we.get('exchange_symbol') or ''}"),
                  ("Market capitalisation (₹ Cr)", we.get("market_cap_inr_cr")), ("Scope 1 (tCO2e)", we.get("scope1_tco2e")),
                  ("Scope 2 (tCO2e)", we.get("scope2_tco2e")), ("Scope 3 (tCO2e)", we.get("scope3_tco2e")),
                  ("Energy (GJ)", we.get("energy_gj")), ("Renewable share (%)", we.get("renewable_share_pct")),
                  ("Water (kL)", we.get("water_kl")), ("Net-zero target", we.get("net_zero_target")),
                  ("Latest BRSR link", we.get("brsr_latest_url"))] if v not in (None, "", [])]
        if rows_:
            _table(doc, ["Item (from web, checked on source page)", "Value"], rows_, widths=[7, 10], font_size=8.5)
        if we.get("_unverified"):
            _note(doc, "Not used - the AI reported these but they could not be found on any source page that names the company: "
                       + "; ".join(f"{k_.replace('_', ' ')} = {v_}" for k_, v_ in we["_unverified"].items()) + ".")
        srcs_ = list(we.get("_verified_sources") or []) if "_verified_sources" in we else \
            [x_.get("url") for x_ in (we.get("sources") or [])] + [c_.get("url") for c_ in (we.get("_citations") or [])]
        srcs_ = [u for u in dict.fromkeys(srcs_) if u][:10]
        if srcs_:
            _note(doc, "Sources: " + "; ".join(srcs_))
    if b_:
        if b_.get("summary"):
            _paras(doc, [b_["summary"]])
        items = [("Scope 1 (tCO2e)", "scope1_tco2e"), ("Scope 2 (tCO2e)", "scope2_tco2e"), ("Scope 3 (tCO2e)", "scope3_tco2e"),
                 ("Energy consumed (GJ)", "energy_total_gj"), ("Renewable energy share (%)", "energy_renewable_pct"),
                 ("Water withdrawal (kL)", "water_withdrawal_kl"), ("Waste generated (t)", "waste_generated_t"),
                 ("GHG assurance", "ghg_assurance"), ("Net-zero target", "net_zero_target"), ("LTIFR", "ltifr_employees"), ("Fatalities", "fatalities")]
        rows = [[lab, f"{b_.get(k):,.2f}" if isinstance(b_.get(k), (int, float)) else str(b_.get(k))] for lab, k in items if b_.get(k) not in (None, "")]
        if rows:
            _table(doc, ["BRSR indicator", "Reported"], rows, widths=[8, 8], font_size=9)
        if b_.get("esg_red_flags"):
            _blt(doc, [f"[{x_.get('severity')}] {x_.get('title')}: {x_.get('evidence')}" for x_ in b_["esg_red_flags"]], 9.5)
    else:
        _paras(doc, ["The BRSR filing was not provided; ESG indicators have not been assessed."])
    pc = bundle.get("pcaf") or {}
    if pc:
        af = pc.get("attribution_factor")
        _paras(doc, [f"Financed emissions are computed under the PCAF Global GHG Accounting and Reporting Standard. The attribution factor "
                     f"is {af:.4%} ({pc.get('denominator_name')}); applied to company Scope 1+2 emissions of "
                     f"{pc.get('company_scope12') or 0:,.0f} tCO2e this gives financed emissions of {pc.get('financed_scope12') or 0:,.0f} tCO2e, "
                     f"with a PCAF data-quality score of {pc.get('data_quality_score')} (1 = verified reported data, 5 = estimated)."
                     if af is not None else "PCAF inputs were incomplete."])
        if pc.get("warning"):
            _note(doc, pc["warning"])

    # ------------------------------------------------------------ risk matrix, covenants, checklist
    doc.add_heading("11. Risk Assessment Matrix", 1)
    rm = AN.risk_matrix(bundle)
    _table(doc, ["Risk area", "Level", "Basis"], [([a, l, b2], {1: {"High": "FEE2E2", "Medium": "FEF3C7", "Low": "DCFCE7"}.get(l)}) for a, l, b2 in rm],
           widths=[5, 2.5, 9.5], font_size=9)
    kr = narrative.get("key_risks") or []
    if kr:
        _p(doc, "Key risks and mitigants", 10.5, bold=True, space_after=2)
        _table(doc, ["Risk", "Mitigant / condition"], [[k_.get("risk"), k_.get("mitigant")] for k_ in kr], widths=[8, 9], font_size=8.5)

    doc.add_heading("12. Proposed Terms, Security, Pricing and Covenants", 1)
    RS.security_pricing_section(doc, bundle, P, num="12")
    doc.add_heading("12.4 Financial covenants (tested on audited annual accounts)", 2)
    _table(doc, ["Covenant", "Threshold", f"Actual {y0}", "Status"], AN.covenants(bundle), widths=[5, 4, 3, 5], font_size=9)
    doc.add_heading("12.5 Information and conduct covenants", 2)
    _blt(doc, (narrative.get("suggested_covenants") if narrative.get("_source") == "AI" else None) or [
        "Quarterly stock and book-debt statements within 20 days of quarter end; drawing power on eligible receivables (ageing < 90 days) only.",
        "Annual stock and receivables audit by a bank-empanelled auditor.",
        "No loans, advances, investments or guarantees to group or related entities without prior written consent.",
        "End-use certificate from the statutory auditor for term-loan disbursements.",
        "Routing of all business receipts through the lending bank's accounts.",
        "Immediate intimation of any auditor qualification, rating action, litigation or regulatory notice."], 9.5)

    doc.add_heading("13. Due-diligence Checklist", 1)
    standard = [
        "KYC of the entity, promoters and guarantors; MCA master data and charge register (CHG) search.",
        "CIBIL / CRILC commercial report and RBI wilful-defaulter / SMA status of the borrower and group entities.",
        "GST returns (GSTR-1 / 3B) for the last 12 months reconciled to reported revenue.",
        "Income-tax returns and tax audit report (Form 3CD) for the years analysed.",
        "Bank statements of all operative accounts for 12 months - turnover, inward/outward cheque returns, related-party transfers, cash withdrawals.",
        "Debtor and creditor ageing with top-10 parties; confirmation of related-party balances.",
        "Plant / office inspection report and stock verification.",
        "Valuation and legal title reports for primary and collateral security.",
    ]
    acts = list(dict.fromkeys([f["action"] for f in sorted(fz["flags"], key=lambda z: {"High": 0, "Medium": 1, "Low": 2}.get(z["severity"], 3))]
                              + (narrative.get("due_diligence_checklist") or []) + standard))
    _table(doc, ["#", "Item to obtain / verify", "Status", "Remarks"], [[str(i), a, "☐ Pending", ""] for i, a in enumerate(acts[:30], 1)],
           widths=[0.8, 10.5, 2.4, 3.3], font_size=8.5)

    # ------------------------------------------------------------ annexures
    doc.add_page_break()
    doc.add_heading("Annexure A - Financial Statements (standardised spread)", 1)
    _table(doc, ["Balance sheet (₹ Cr)"] + cols, bs_rows, font_size=8)
    pl_rows = [[label] + [_fmtv(d[y].get(k)) for y in cols] for k, label, _, _ in schema.PROFIT_LOSS + schema.CASH_FLOW
               if any(d[y].get(k) not in (None, 0) for y in cols)]
    _p(doc, "", space_after=4)
    _table(doc, ["Profit & loss and cash flow (₹ Cr)"] + cols, pl_rows, font_size=8)
    _glossary(doc)
    doc.add_heading("Annexure C - Sources, Method and Processing Cost", 1)
    _blt(doc, [
        f"Financial statements: {bundle.get('source_file') or 'uploaded documents'} (extraction: {bundle.get('method')}).",
        f"BRSR: {(bundle.get('brsr') or {}).get('method', 'not provided')}.",
        f"News: {len(news)} items (Google News RSS and configured sources).",
        f"AI models used: {meta.get('model', 'none (rule-based only)')}.",
        (f"Processing cost of this memorandum: ₹{meta['ai_cost']['inr']:,.2f} (USD {meta['ai_cost']['usd']:,.4f}; "
         f"{meta['ai_cost']['calls']} AI calls, {meta['ai_cost']['cached']} served from cache; basis: {meta['ai_cost'].get('basis', 'estimate')}; "
         f"USD/INR {meta.get('fx', 88)})."
         if meta.get("ai_cost") else "Processing cost: ₹0 (rule-based only)."),
        "Ratios, scores and forensic tests are computed deterministically; AI is used only to read documents and draft narrative.",
        "Benchmarks are generic bank norms; apply the bank's credit policy and sector-specific norms.",
    ], 9)
    _note(doc, "Disclaimer: This memorandum is decision support and does not constitute a sanction. Automated extraction can misread "
               "scanned documents; material figures must be verified against the audited financial statements.")
    chk_ = bundle.get("checks") or []
    if chk_:
        doc.add_page_break()
        doc.add_heading("Annexure D - Verification Log", 1)
        _paras(doc, ["Every figure in this memorandum was put through automated logic checks before the report was written: the statements must "
                     "add up, key ratios are recomputed independently, implausible year-on-year jumps are flagged, the rating and recommendation "
                     "must agree, every figure in AI-written text must trace back to the financial data, and anything taken from the web is used "
                     "only if the source page could be opened and shows the company and the figure."])
        ig_ = bundle.get("integrity") or {}
        if ig_:
            _keyvalue(doc, [["Case status", ig_["status"].replace("_", " ")],
                            ["Final recommendation allowed", "Yes" if ig_["final_recommendation_allowed"] else "No"],
                            ["Required lines not read", ", ".join(ig_["missing_required_fields"]) or "None"],
                            ["Decision metrics not computed", ", ".join(ig_["blocked_metrics"]) or "None"],
                            ["Failed checks", "; ".join(ig_["failed_checks"]) or "None"],
                            ["Warnings", str(ig_.get("warnings", 0))]], widths=(5, 12))
        _table(doc, ["Area", "Check", "Result", "Detail"],
               [([c_["area"], c_["check"], c_["result"], c_["detail"]], {2: {"PASS": "DCFCE7", "WARN": "FEF3C7", "FAIL": "FEE2E2"}.get(c_["result"])})
                for c_ in chk_], widths=[2.6, 6.2, 1.6, 6.6], font_size=8)
        fch_ = (bundle.get("meta") or {}).get("figure_changes") or []
        _p(doc, "Figures changed by hand after reading", 10, bold=True, space_after=2)
        if fch_:
            _table(doc, ["When", "By", "Line", "Year", "Was", "Now", "Reason"],
                   [[c_["ts"].replace("T", " ")[:16], c_["user"], c_["line"], c_["year"], _fmtv(c_["old"]), _fmtv(c_["new"]), c_["reason"][:160]]
                    for c_ in fch_[:60]], widths=[2.2, 1.6, 3.4, 1.3, 1.6, 1.6, 5.3], font_size=7.5)
        else:
            _paras(doc, ["None - every figure is as read from the documents (and the second reading, where it ran)."])
        rem_ = (narrative or {}).get("_removed_sentences") or []
        if rem_:
            _p(doc, "AI-written sentences removed (figures could not be traced to the data)", 10, bold=True, space_after=2)
            _table(doc, ["Section", "Removed sentence", "Untraced figure(s)"],
                   [[r_["where"].replace("_", " "), r_["sentence"][:300], r_["untraced"]] for r_ in rem_[:30]], widths=[3, 11, 3], font_size=7.5)
    _static_toc(doc, toc_anchor)
    buf_ = io.BytesIO()
    doc.save(buf_)
    return buf_.getvalue()


def _static_toc(doc, anchor):
    """Contents list written from the actual headings (no Word field to update by hand)."""
    import copy
    heads = [(p_.style.name, p_.text.strip()) for p_ in doc.paragraphs
             if p_.style is not None and p_.style.name in ("Heading 1", "Heading 2") and p_.text.strip()]
    prev = anchor
    for lvl, txt in heads:
        np_ = copy.deepcopy(anchor._p)
        prev._p.addnext(np_)
        from docx.text.paragraph import Paragraph
        para = Paragraph(np_, anchor._parent)
        r = para.add_run(txt)
        r.font.size = Pt(10.5 if lvl == "Heading 1" else 9.5)
        r.bold = lvl == "Heading 1"
        r.font.color.rgb = INK if lvl == "Heading 1" else GREY
        para.paragraph_format.left_indent = Cm(0 if lvl == "Heading 1" else 0.8)
        para.paragraph_format.space_after = Pt(3 if lvl == "Heading 1" else 1)
        para.paragraph_format.space_before = Pt(5 if lvl == "Heading 1" else 0)
        prev = para
    anchor._p.getparent().remove(anchor._p)


def _glossary(doc):
    doc.add_heading("Annexure B - Ratio Definitions", 1)
    rows = []
    for key, name, cat, _t, _f, direction, good, bad, _n in R.RATIOS:
        what, how = EXPL.RATIO_HELP.get(key, ("", ""))
        bench = ""
        if good is not None:
            bench = ("≥ " if direction == "high" else "≤ " if direction == "low" else "") + R.fmt(key, good)
        rows.append([name, what, how, bench])
    _table(doc, ["Ratio", "Definition", "Interpretation", "Benchmark"], rows, widths=[3.6, 5.0, 6.6, 1.8], font_size=7.5)
    _p(doc, "Scores", 10.5, bold=True, space_after=2)
    _table(doc, ["Score", "Meaning"], [[a, b] for a, b in EXPL.SCORE_HELP], widths=[4.5, 12.5], font_size=8)
