"""Excel workbook with live formulas (3 core tabs + analysis tabs from report_excel_extra).
 Tab 1  Balance Sheet            - full statements (BS, P&L, Cash flow, disclosures) + derived rows (formulas)
 Tab 2  Ratios                   - every ratio as a live formula on Tab 1, benchmark + traffic-light formatting
 Tab 3  PCAF Financed Emissions  - attribution factor, financed emissions, data-quality score (formulas)
"""
import io
import re

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import ratios as R
from . import schema

FONT = "Arial"
NAVY = "1C1917"  # ink (house style)
F_HEAD = Font(name=FONT, bold=True, color="FFFFFF", size=10)
F_SEC = Font(name=FONT, bold=True, color=NAVY, size=10)
F_IN = Font(name=FONT, color="0000FF", size=10)       # hard-coded inputs
F_CALC = Font(name=FONT, color="000000", size=10)     # formulas
F_LINK = Font(name=FONT, color="008000", size=10)     # cross-sheet links
F_BOLD = Font(name=FONT, bold=True, size=10)
F_NOTE = Font(name=FONT, italic=True, color="595959", size=9)
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_SEC = PatternFill("solid", fgColor="FBE3D3")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_GOOD = PatternFill("solid", fgColor="C6EFCE")
FILL_WATCH = PatternFill("solid", fgColor="FFEB9C")
FILL_BAD = PatternFill("solid", fgColor="FFC7CE")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
NUM = '#,##0.00;(#,##0.00);"-"'
FMT = {"x": '0.00"x";(0.00"x");"-"', "%": '0.0%;(0.0%);"-"', "days": '#,##0" d";(#,##0" d");"-"', "cr": NUM}
GUARDS = {"cash_yield": ["interest_income"], "rpt_rev": ["rpt_sales", "rpt_purchases"],
          "icr": ["finance_cost"], "dscr": ["current_maturities_ltd"], "ebitda_int": ["finance_cost"], "implied_rate": ["finance_cost"]}

# same "not meaningful" rules as core/ratios.compute, written as live formulas so the workbook agrees with the memorandum
NM_GUARDS = {
    "icr": ("ABS({finance_cost})<MAX(0.005,0.0001*ABS({revenue}))", "n.a. - nil interest"),
    "ebitda_int": ("ABS({finance_cost})<MAX(0.005,0.0001*ABS({revenue}))", "n.a. - nil interest"),
    "implied_rate": ("ABS({finance_cost})<MAX(0.005,0.0001*ABS({revenue}))", "n.a. - nil interest"),
    "debt_ebitda": ("{ebitda}<=0", "n.m. - EBITDA negative"),
    "netdebt_ebitda": ("{ebitda}<=0", "n.m. - EBITDA negative"),
    "cfo_ebitda": ("{ebitda}<=0", "n.m. - EBITDA negative"),
    "cfo_pat": ("{pat}<=0", "n.m. - loss year"),
    "etr": ("{pbt}<=0", "n.m. - loss year"),
    "other_inc_pbt": ("{pbt}<=0", "n.m. - loss year"),
    "inventory_days": ("AND({revenue}<>0,{cogs}<0.05*ABS({revenue}))", "n.m. - service business"),
    "creditor_days": ("AND({revenue}<>0,{cogs}<0.05*ABS({revenue}))", "n.m. - service business"),
    "ccc": ("AND({revenue}<>0,{cogs}<0.05*ABS({revenue}))", "n.m. - service business"),
    "rev_growth": ("{revenue@p}<=0", "n.m. - nil / negative base"),
    "ebitda_growth": ("{ebitda@p}<=0", "n.m. - nil / negative base"),
    "pat_growth": ("{pat@p}<=0", "n.m. - nil / negative base"),
    "debt_growth": ("{total_debt@p}<=0", "n.m. - nil / negative base"),
}

S1, S2, S3 = "Balance Sheet", "Ratios", "PCAF Financed Emissions"


def _link(ref):
    """Link to another sheet that stays blank-aware: Excel turns a link to an empty cell into 0, which would read as 'nil'."""
    return f'=IF(ISBLANK({ref}),"n.a. - not read",{ref})'


def _first(*vals):
    """First value that is known (0 counts as known); None when none is - so 'not reported' never becomes 0."""
    return next((v for v in vals if v is not None and v != ""), None)


def _q(sheet):
    return f"'{sheet}'"


def _header(ws, row, values, widths=None):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, Alignment(horizontal="center", vertical="center", wrap_text=True), BOX
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def build(bundle, pcaf_inputs=None):
    years = bundle["years"]  # latest first
    cols_years = list(reversed(years))  # oldest -> latest left to right (credit convention)
    data = bundle["data"]
    company = bundle.get("company") or "Company"
    wb = Workbook()

    # ------------------------------------------------------------------ TAB 1
    ws = wb.active
    ws.title = S1
    ws["A1"] = f"{company} - Financial Statements ({(bundle.get('statement_type') or 'standalone').title()})"
    ws["A1"].font = Font(name=FONT, bold=True, size=13, color=NAVY)
    ws["A2"] = "All figures INR Crore. Blue = extracted from the uploaded PDF (edit if needed); black = formulas. Source: " + (bundle.get("source_file") or "uploaded annual report")
    ws["A2"].font = F_NOTE
    hdr_row = 4
    _header(ws, hdr_row, ["Key", "Particulars"] + cols_years, [24, 52] + [16] * len(cols_years))
    ycol = {y: get_column_letter(3 + i) for i, y in enumerate(cols_years)}
    row_of = {}
    r = hdr_row + 1
    groups = [("BALANCE SHEET", schema.BALANCE_SHEET), ("STATEMENT OF PROFIT & LOSS", schema.PROFIT_LOSS),
              ("CASH FLOW STATEMENT", schema.CASH_FLOW), ("KEY DISCLOSURES (Notes)", schema.DISCLOSURES)]
    for title, items in groups:
        ws.cell(row=r, column=1, value=title).font = F_SEC
        for c in range(1, 3 + len(cols_years)):
            ws.cell(row=r, column=c).fill = FILL_SEC
        r += 1
        last_sec = None
        for key, label, sec, _ in items:
            if sec != last_sec and title == "BALANCE SHEET":
                ws.cell(row=r, column=2, value=sec).font = Font(name=FONT, bold=True, italic=True, size=10)
                r += 1
                last_sec = sec
            ws.cell(row=r, column=1, value=key).font = F_NOTE
            lc = ws.cell(row=r, column=2, value=label)
            lc.font = F_BOLD if key.startswith("total") else F_CALC
            for y in cols_years:
                v = data[y].get(key)
                c = ws[f"{ycol[y]}{r}"]
                c.value = v
                c.font = F_IN
                c.number_format = NUM if key != "promoter_pledge_pct" else '0.00"%"'
                c.border = BOX
            row_of[key] = r
            r += 1
        r += 1
    # Check row
    ws.cell(row=r, column=2, value="Check: Total Assets - Total Equity & Liabilities (should be 0)").font = F_NOTE
    for y in cols_years:
        c = ws[f"{ycol[y]}{r}"]
        c.value = f"={ycol[y]}{row_of['total_assets']}-{ycol[y]}{row_of['total_equity_and_liabilities']}"
        c.number_format, c.font = NUM, F_CALC
    r += 2
    ws.cell(row=r, column=1, value="DERIVED FIGURES (formulas)").font = F_SEC
    for c in range(1, 3 + len(cols_years)):
        ws.cell(row=r, column=c).fill = FILL_SEC
    r += 1
    derived_rows = {}
    for key, label, _t in R.DERIVED:
        derived_rows[key] = r
        r += 1
    row_of.update(derived_rows)

    def cell_local(key, col):
        return f"{col}{row_of[key]}"

    for key, label, t in R.DERIVED:
        rr = derived_rows[key]
        ws.cell(row=rr, column=1, value=key).font = F_NOTE
        ws.cell(row=rr, column=2, value=label).font = F_BOLD
        for y in cols_years:
            f = R.excel_formula(t, cell_local, ycol[y], None).replace('IFERROR(', '', 1)
            f = "=" + f[1:].rsplit(',"n.a.")', 1)[0]
            blk = sorted({k for k, _p in R.base_keys(t) if k in R.BLOCKING and k in row_of})
            if blk:
                f = '=IF(OR(' + ",".join(f"ISBLANK({cell_local(k, ycol[y])})" for k in blk) + '),"n.a. - input missing",' + f[1:] + ")"
            c = ws[f"{ycol[y]}{rr}"]
            c.value, c.font, c.number_format, c.border = f, F_CALC, NUM, BOX
    ws.freeze_panes = ws[f"C{hdr_row + 1}"]

    # ------------------------------------------------------------------ TAB 2
    w2 = wb.create_sheet(S2)
    w2["A1"] = f"{company} - Ratio Analysis (live formulas on '{S1}')"
    w2["A1"].font = Font(name=FONT, bold=True, size=13, color=NAVY)
    w2["A2"] = "Green = within benchmark, amber = watch, red = outside benchmark. Benchmarks are general bank norms; adjust to your policy."
    w2["A2"].font = F_NOTE
    _header(w2, 4, ["Category", "Ratio"] + cols_years + ["Good if", "Benchmark (good)", "Red line", "Formula / note"],
            [18, 42] + [14] * len(cols_years) + [10, 14, 12, 60])
    ycol2 = {y: get_column_letter(3 + i) for i, y in enumerate(cols_years)}

    def cell_xs(key, col):
        return f"{_q(S1)}!{col}{row_of[key]}"

    rr = 5
    last_cat = None
    ratio_rows = {}
    for key, name, cat, t, f, direction, good, bad, note in R.RATIOS:
        if cat != last_cat:
            w2.cell(row=rr, column=1, value=cat).font = F_SEC
            for c in range(1, 7 + len(cols_years)):
                w2.cell(row=rr, column=c).fill = FILL_SEC
            rr += 1
            last_cat = cat
        w2.cell(row=rr, column=2, value=name).font = F_CALC
        for i, y in enumerate(cols_years):
            prev_col = ycol[cols_years[i - 1]] if i > 0 else None
            formula = R.excel_formula(t, cell_xs, ycol[y], prev_col)
            if key in GUARDS and formula != '="n.a."':
                cond = "+".join(f"ABS({cell_xs(k, ycol[y])})" for k in GUARDS[key])
                formula = f'=IF({cond}=0,"n.a.",{formula[1:]})'
            blk = [(k, pp) for k, pp in R.base_keys(t) if k in R.BLOCKING]
            if blk and formula != '="n.a."':
                conds = [f"ISBLANK({cell_xs(k, prev_col if pp else ycol[y])})" for k, pp in sorted(set(blk)) if (prev_col or not pp)]
                if conds:
                    formula = '=IF(OR(' + ",".join(conds) + '),"n.a. - input missing",' + formula[1:] + ")"
            if key in NM_GUARDS and formula != '="n.a."':
                tmpl, lab = NM_GUARDS[key]
                cond = re.sub(r"\{(\w+)(@p)?\}", lambda m: cell_xs(m.group(1), prev_col if m.group(2) else ycol[y]), tmpl)
                formula = f'=IF({cond},"{lab}",{formula[1:]})'
            c = w2[f"{ycol2[y]}{rr}"]
            c.value, c.number_format, c.font, c.border = formula, FMT[f], F_LINK, BOX
            c.alignment = Alignment(horizontal="right")
        nc = 3 + len(cols_years)
        w2.cell(row=rr, column=nc, value={"high": "Higher", "low": "Lower", "band": "Within band"}.get(direction, "-")).font = F_NOTE
        if good is not None:
            w2.cell(row=rr, column=nc + 1, value=good).number_format = FMT[f]
            w2.cell(row=rr, column=nc + 2, value=bad).number_format = FMT[f]
            w2.cell(row=rr, column=nc + 1).font = F_IN
            w2.cell(row=rr, column=nc + 2).font = F_IN
        w2.cell(row=rr, column=nc + 3, value=(note + "  | " if note else "") + t.replace("{", "").replace("}", "").replace("@p", "(prev yr)")).font = F_NOTE
        # traffic lights (formula-driven so they react to edits)
        if direction in ("high", "low") and good is not None:
            rng = f"{ycol2[cols_years[0]]}{rr}:{ycol2[cols_years[-1]]}{rr}"
            gcell, bcell = f"${get_column_letter(nc + 1)}${rr}", f"${get_column_letter(nc + 2)}${rr}"
            first = f"{ycol2[cols_years[0]]}{rr}"
            if direction == "high":
                rules = [(f"AND(ISNUMBER({first}),{first}>={gcell})", FILL_GOOD), (f"AND(ISNUMBER({first}),{first}<{bcell})", FILL_BAD),
                         (f"ISNUMBER({first})", FILL_WATCH)]
            else:
                rules = [(f"AND(ISNUMBER({first}),{first}<={gcell})", FILL_GOOD), (f"AND(ISNUMBER({first}),{first}>{bcell})", FILL_BAD),
                         (f"ISNUMBER({first})", FILL_WATCH)]
            for formula, fill in rules:
                w2.conditional_formatting.add(rng, FormulaRule(formula=[formula], fill=fill, stopIfTrue=True))
        ratio_rows[key] = rr
        rr += 1
    # model scores (computed in the app)
    rr += 1
    w2.cell(row=rr, column=1, value="Forensic & distress scores").font = F_SEC
    for c in range(1, 7 + len(cols_years)):
        w2.cell(row=rr, column=c).fill = FILL_SEC
    rr += 1
    fz = bundle["forensics"]
    latest_col = ycol2[cols_years[-1]]
    for label, val, note in [
        ("Beneish M-score (8-variable)", fz.get("beneish"), "> -1.78 = likely manipulator (computed by app, latest vs previous year)"),
        ("Piotroski F-score (0-9)", fz.get("piotroski"), ">= 7 strong, <= 3 weak (computed by app)"),
        ("Internal credit score (0-100)", fz.get("score"), fz.get("grade")),
        ("Red flags raised (High / Medium / Low)", "{} / {} / {}".format(*[sum(1 for f in fz["flags"] if f["severity"] == s) for s in ("High", "Medium", "Low")]), "See Word report for detail"),
    ]:
        w2.cell(row=rr, column=2, value=label).font = F_BOLD
        c = w2[f"{latest_col}{rr}"]
        c.value, c.font = val, F_IN
        w2.cell(row=rr, column=6 + len(cols_years), value=note).font = F_NOTE
        rr += 1
    w2.freeze_panes = "C5"

    # ------------------------------------------------------------------ TAB 3
    w3 = wb.create_sheet(S3)
    p = pcaf_inputs or {}
    brsr = bundle.get("brsr") or {}
    latest = ycol[cols_years[-1]]
    w3["A1"] = f"{company} - PCAF Financed Emissions (Business loans / Listed equity & bonds)"
    w3["A1"].font = Font(name=FONT, bold=True, size=13, color=NAVY)
    w3["A2"] = "Yellow cells are inputs - change them and every result recalculates. Green = linked from the Balance Sheet tab (latest year)."
    w3["A2"].font = F_NOTE
    _header(w3, 4, ["Item", "Value", "Unit", "Source / note"], [48, 18, 16, 80])
    rows = [
        ("INPUTS", None, None, None, "sec"),
        ("Bank's outstanding exposure", p.get("outstanding_cr"), "INR Cr", "Funded outstanding as on reporting date (user input; blank = not entered)", "in"),
        ("Asset class", p.get("asset_class", "Business loans / unlisted equity"), "", "'Listed equity & corporate bonds' uses EVIC; else equity + debt", "in"),
        ("Market capitalisation (listed only)", p.get("market_cap"), "INR Cr", "Year-end market cap (user input / exchange data; blank = not known)", "in"),
        ("Total equity (book)", _link(f"{_q(S1)}!{latest}{row_of['total_equity']}"), "INR Cr", f"Linked: {S1} latest year", "link"),
        ("Total debt (book)", _link(f"{_q(S1)}!{latest}{row_of['total_debt']}"), "INR Cr", f"Linked: {S1} derived Total Debt", "link"),
        ("Non-controlling interest", _link(f"{_q(S1)}!{latest}{row_of['minority_interest']}"), "INR Cr", f"Linked: {S1}", "link"),
        ("Revenue", _link(f"{_q(S1)}!{latest}{row_of['revenue']}"), "INR Cr", f"Linked: {S1}", "link"),
        ("Scope 1 emissions (reported)", _first(p.get("scope1"), brsr.get("scope1_tco2e")), "tCO2e", "BRSR Principle 6 - leave BLANK if not reported (0 means reported as zero)", "in"),
        ("Scope 2 emissions (reported)", _first(p.get("scope2"), brsr.get("scope2_tco2e")), "tCO2e", "BRSR Principle 6 - leave BLANK if not reported", "in"),
        ("Scope 3 emissions (reported, optional)", _first(p.get("scope3"), brsr.get("scope3_tco2e")), "tCO2e", "BRSR Core / voluntary - leave BLANK if not reported", "in"),
        ("Emissions third-party assured?", "Yes" if p.get("verified") else "No", "Yes/No", "Reasonable/limited assurance on GHG -> PCAF score 1, else 2", "in"),
        ("Sector emission factor (fallback)", p.get("sector_ef") or None, "tCO2e / INR Cr revenue", f"Used only when no reported emissions. Sector: {p.get('sector_name') or '-'} (ILLUSTRATIVE default - replace with approved factor)", "in"),
        ("CALCULATION", None, None, None, "sec"),
        # unknown stays unknown: a blank input gives "n.a." - never a zero that would read as 'no emissions'
        ("Attribution denominator", "=IF(AND(LEFT(B7,6)=\"Listed\",ISNUMBER(B8),B8>0),IF(ISNUMBER(B10),B8+B10+N(B11),\"n.a. - debt not read\"),"
                                    "IF(AND(ISNUMBER(B9),ISNUMBER(B10)),B9+B10,\"n.a. - equity or debt not read\"))", "INR Cr", "EVIC for listed; equity + debt for unlisted / business loans", "f"),
        ("Attribution factor", "=IF(AND(ISNUMBER(B6),ISNUMBER(B19)),IF(B19>0,MIN(1,B6/B19),\"n.a.\"),\"n.a. - exposure or denominator missing\")", "%",
         "Outstanding / denominator (capped at 100%)", "f"),
        ("Company Scope 1+2 used", "=IF(AND(ISNUMBER(B13),ISNUMBER(B14)),B13+B14,IF(AND(ISNUMBER(B17),ISNUMBER(B12)),B17*B12,\"n.a. - not reported\"))", "tCO2e",
         "Reported Scope 1 AND 2 if both available, else revenue x approved sector factor; a single reported scope is not used as the total", "f"),
        ("Financed emissions - Scope 1+2", "=IF(AND(ISNUMBER(B20),ISNUMBER(B21)),B20*B21,\"n.a.\")", "tCO2e", "Attribution factor x company Scope 1+2", "f"),
        ("Financed emissions - Scope 3", "=IF(AND(ISNUMBER(B20),ISNUMBER(B15)),B20*B15,\"n.a. - not reported\")", "tCO2e", "Reported separately per PCAF", "f"),
        ("PCAF data quality score", "=IF(AND(ISNUMBER(B13),ISNUMBER(B14)),IF(B16=\"Yes\",1,2),IF(AND(ISNUMBER(B17),ISNUMBER(B12)),4,\"n.a.\"))", "1 (best) - 5",
         "1 verified, 2 unverified, 3 physical activity, 4 revenue-based, 5 asset-based", "f"),
        ("Financed emission intensity", "=IF(AND(ISNUMBER(B22),ISNUMBER(B6),N(B6)>0),B22/B6,\"n.a.\")", "tCO2e / INR Cr lent", "Carbon footprint of this exposure", "f"),
        ("Company economic emission intensity", "=IF(AND(ISNUMBER(B21),ISNUMBER(B12),N(B12)>0),B21/B12,\"n.a.\")", "tCO2e / INR Cr revenue", "Compare with sector peers", "f"),
    ]
    rr = 5
    for label, val, unit, note, kind in rows:
        if kind == "sec":
            w3.cell(row=rr, column=1, value=label).font = F_SEC
            for c in range(1, 5):
                w3.cell(row=rr, column=c).fill = FILL_SEC
            rr += 1
            continue
        w3.cell(row=rr, column=1, value=label).font = F_BOLD if kind == "f" else F_CALC
        c = w3.cell(row=rr, column=2, value=val)
        c.border = BOX
        if kind == "in":
            c.font, c.fill = F_IN, FILL_INPUT
        elif kind == "link":
            c.font = F_LINK
        else:
            c.font = F_BOLD
        if unit in ("INR Cr", "tCO2e") or "tCO2e /" in unit:
            c.number_format = '#,##0.00;(#,##0.00);"-"'
        if unit == "%":
            c.number_format = "0.0000%"
        w3.cell(row=rr, column=3, value=unit).font = F_NOTE
        w3.cell(row=rr, column=4, value=note).font = F_NOTE
        rr += 1
    # sanity: row numbers used in formulas above
    assert w3["A6"].value == "Bank's outstanding exposure" and w3["A19"].value == "Attribution denominator", "PCAF layout shifted"
    rr += 1
    w3.cell(row=rr, column=1, value="Method: PCAF Global GHG Accounting & Reporting Standard, Part A (Financed Emissions), business loans & listed equity/corporate bonds.").font = F_NOTE
    w3.cell(row=rr + 1, column=1, value="Only sector factors marked approved in data/emission_factors.csv (Admin › Bank parameters) are used; unapproved rows are ignored.").font = F_NOTE

    for wsx in (ws, w2, w3):
        wsx.sheet_view.showGridLines = False
        wsx.sheet_properties.pageSetUpPr.fitToPage = True
    try:
        from . import report_excel_extra as XE
        XE.add_tabs(wb, bundle)
    except Exception as e:  # never lose the three core tabs because of an optional tab
        wx = wb.create_sheet("Extra tabs - error")
        wx["A1"] = f"Additional analysis tabs could not be built: {e}"
    chk = bundle.get("checks") or []
    if chk:
        wc = wb.create_sheet("Checks", index=1)
        wc["A1"] = f"{bundle.get('company')} - logic and source checks run before this workbook was written"
        wc["A1"].font = Font(name=FONT, bold=True, size=13, color=NAVY)
        ig = bundle.get("integrity") or {}
        wc["A2"] = (f"Case status: {ig['status'].replace('_', ' ')} - {ig.get('detail', '')}. " if ig else "") + \
            "FAIL = fix the figures before use; WARN = read and confirm; the same log is Annexure D of the memorandum."
        wc["A2"].font = F_NOTE
        for j, h in enumerate(["Area", "Check", "Result", "Detail"], 1):
            c = wc.cell(row=4, column=j, value=h)
            c.font, c.fill = F_HEAD, FILL_HEAD
        fills = {"PASS": FILL_GOOD, "WARN": FILL_WATCH, "FAIL": FILL_BAD}
        for i, c_ in enumerate(chk, 5):
            for j, k in enumerate(["area", "check", "result", "detail"], 1):
                cell = wc.cell(row=i, column=j, value=c_[k])
                cell.font = F_CALC
            wc.cell(row=i, column=3).fill = fills.get(c_["result"], FILL_WATCH)
        fch = (bundle.get("meta") or {}).get("figure_changes") or []
        r0 = 5 + len(chk) + 2
        wc.cell(row=r0, column=1, value="Figures changed after reading (audit trail - entries are never edited or removed)").font = Font(name=FONT, bold=True, size=11, color=NAVY)
        if not fch:
            wc.cell(row=r0 + 1, column=1, value="None - every figure is as read from the documents.").font = F_NOTE
        else:
            for j, h in enumerate(["When", "Line (year)", "By", "Was -> now / reason"], 1):
                c = wc.cell(row=r0 + 1, column=j, value=h)
                c.font, c.fill = F_HEAD, FILL_HEAD
            for i, c_ in enumerate(fch, r0 + 2):
                vals = [str(c_.get("ts", "")).replace("T", " ")[:16], f"{c_.get('line')} ({c_.get('year')})", c_.get("user"),
                        f"{c_.get('old')} -> {c_.get('new')} | {c_.get('reason')}"]
                for j, v in enumerate(vals, 1):
                    wc.cell(row=i, column=j, value=v).font = F_CALC
        for col, w in zip("ABCD", (16, 62, 10, 90)):
            wc.column_dimensions[col].width = w
        wc.sheet_view.showGridLines = False
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
