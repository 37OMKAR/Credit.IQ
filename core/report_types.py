"""Report centre: the different documents a credit officer needs from one analysis.

Every builder takes (bundle, narrative, officer, branch) and returns .docx bytes.
"""
import datetime as dt

from docx.shared import Pt

from . import analyst as AN
from . import credit_tools as CT
from . import ratios as R
from . import report_sections as RS
from . import report_word as W
from .schema import g


def _key_fin_rows(b):
    years = list(reversed(b["years"]))
    d, r = b["data"], b["ratios"]
    lines = [("Revenue", lambda y: W._fmtv(g(d[y], "revenue"))), ("EBITDA", lambda y: W._fmtv(g(d[y], "ebitda"))),
             ("EBITDA margin", lambda y: R.fmt("ebitda_margin", r[y].get("ebitda_margin"))), ("PAT", lambda y: W._fmtv(g(d[y], "pat"))),
             ("Net worth", lambda y: W._fmtv(g(d[y], "total_equity"))), ("Total debt", lambda y: W._fmtv(g(d[y], "total_debt"))),
             ("TOL / TNW", lambda y: R.fmt("tol_tnw", r[y].get("tol_tnw"))), ("Current ratio", lambda y: R.fmt("current_ratio", r[y].get("current_ratio"))),
             ("Interest coverage", lambda y: R.fmt("icr", r[y].get("icr"))), ("DSCR", lambda y: R.fmt("dscr", r[y].get("dscr"))),
             ("Debtor days", lambda y: R.fmt("debtor_days", r[y].get("debtor_days"))), ("CFO", lambda y: W._fmtv(d[y].get("cfo")))]
    return years, [[lab] + [f(y) for y in years] for lab, f in lines]


def _signoff(doc):
    W._p(doc, "", space_after=6)
    W._table(doc, ["", "Name / designation", "Signature", "Date"],
             [["Prepared by", "", "", ""], ["Reviewed by", "", "", ""], ["Recommended by", "", "", ""], ["Sanctioning authority", "", "", ""]],
             widths=[3.6, 6, 4, 3.4], font_size=9)


def _flags_sorted(b):
    return sorted(b["forensics"]["flags"], key=lambda f: {"High": 0, "Medium": 1, "Low": 2}.get(f["severity"], 3))


# ------------------------------------------------------------------ 1. credit committee note
def committee_note(b, narr, officer="", branch=""):
    p = CT.proposal(b)
    fz = b["forensics"]
    doc = RS.new_doc("Credit Committee Note", "Summary for the sanctioning committee - read with the full Credit Appraisal Memorandum", b, officer, branch)
    doc.add_heading("1. Proposal", 1)
    RS.proposal_snapshot(doc, b, p)
    doc.add_heading("2. Rating and recommendation", 1)
    W._keyvalue(doc, [["Internal rating", f"{fz['grade']} (score {fz['score']:.0f}/100)"],
                      ["Recommendation", narr.get("recommendation", "-")],
                      ["Red flags", f"{sum(1 for f in fz['flags'] if f['severity'] == 'High')} high / {len(fz['flags'])} total"]])
    W._paras(doc, [narr.get("recommendation_rationale") or AN.rating_rationale(b, narr)])
    doc.add_heading("3. Key financials (₹ Cr)", 1)
    yrs, rows = _key_fin_rows(b)
    W._table(doc, ["Particulars"] + yrs, rows, font_size=8.5)
    s_, c_ = AN.strengths_concerns(b)
    doc.add_heading("4. Strengths and concerns", 1)
    n = max(len(s_[:5]), len(c_[:5]))
    W._table(doc, ["Strengths", "Concerns"], [[s_[i] if i < len(s_[:5]) else "", c_[i] if i < len(c_[:5]) else ""] for i in range(n)] or [["-", "-"]],
             widths=[8.5, 8.5], font_size=8.5)
    doc.add_heading("5. Repayment capacity", 1)
    rows_, _ = CT.projections(b, p)
    sm = CT.dscr_summary(rows_)
    sc = CT.scenarios(b, p)
    be = CT.breakeven(b, p)
    W._paras(doc, [f"Projected average DSCR {RS.xr(sm['avg'])}, minimum {RS.xr(sm['min'])}. Downside-case minimum DSCR "
                   f"{RS.xr(sc[2]['min_dscr'])}; EBITDA can fall about {(be or 0):.0%} before minimum DSCR breaches 1.00x."])
    W._table(doc, ["Scenario", "Avg DSCR", "Min DSCR", "Y+1 ICR"], [[s["scenario"], RS.xr(s["avg_dscr"]), RS.xr(s["min_dscr"]), RS.xr(s["y1_icr"])] for s in sc],
             widths=[5, 4, 4, 4], font_size=8.5)
    doc.add_heading("6. Key risks and mitigants", 1)
    kr = narr.get("key_risks") or []
    if kr:
        W._table(doc, ["Risk", "Mitigant / condition"], [[k.get("risk"), k.get("mitigant")] for k in kr[:6]], widths=[8, 9], font_size=8.5)
    else:
        W._table(doc, ["Risk area", "Level", "Basis"], [r_ for r_ in AN.risk_matrix(b) if r_[1] in ("High", "Medium")] or [["-", "-", "-"]],
                 widths=[5, 2.5, 9.5], font_size=8.5)
    ews, nt = CT.rbi_ews(b, p)
    doc.add_heading("7. Early-warning signals", 1)
    W._paras(doc, [f"{nt} of {len(ews)} RBI early-warning signals triggered. {CT.ews_action(nt)}"])
    trig = [e for e in ews if e[1] == "Triggered"]
    if trig:
        W._blt(doc, [f"{e[0]} - {e[2]}" for e in trig], 9)
    doc.add_heading("8. Key conditions", 1)
    W._table(doc, ["Covenant", "Threshold", "Actual", "Status"], AN.covenants(b), widths=[5, 4, 3, 5], font_size=8.5)
    _signoff(doc)
    return RS.save(doc)


# ------------------------------------------------------------------ 2. annual review / renewal
def annual_review(b, narr, officer="", branch=""):
    p = CT.proposal(b)
    fz = b["forensics"]
    years = b["years"]
    doc = RS.new_doc("Annual Review / Renewal Note", "Review of the borrower's performance, account conduct and covenant compliance", b, officer, branch)
    ews, nt = CT.rbi_ews(b, p)
    cov = AN.covenants(b)
    breaches = [c for c in cov if c[3].startswith("Currently breached")]
    grade = CT.grade_code(fz["grade"])
    prev = (p.get("prev_grade") or "").split(" ")[0]
    order = ["CI-1", "CI-2", "CI-3", "CI-4", "CI-5", "CI-6", "CI-7"]
    migration = "-"
    if prev in order and grade in order:
        dlt = order.index(grade) - order.index(prev)
        migration = "Stable" if dlt == 0 else (f"Downgrade by {dlt} notch(es)" if dlt > 0 else f"Upgrade by {-dlt} notch(es)")
    if nt >= 4 or grade in ("CI-6", "CI-7") or int(p.get("dpd") or 0) > 60:
        rec = "Reduce exposure / place under exit or resolution watch"
    elif breaches or nt >= 2 or "Downgrade" in migration:
        rec = "Renew with tightened conditions and closer monitoring"
    else:
        rec = "Renew at existing terms"
    doc.add_heading("1. Review summary", 1)
    W._keyvalue(doc, [["Facilities", "; ".join(x for x in [f"TL ₹{p['tl_amount']:,.2f} Cr" if p.get("tl_amount") else "",
                                                             f"CC ₹{p['cc_amount']:,.2f} Cr" if p.get("cc_amount") else "",
                                                             f"NFB ₹{p['nfb_amount']:,.2f} Cr" if p.get("nfb_amount") else ""] if x) or (p.get("facility") or "-")],
                      ["Account status (IRACP)", CT.sma_status(p.get("dpd"))],
                      ["Internal rating", f"{fz['grade']} (score {fz['score']:.0f}/100)"],
                      ["Previous rating / migration", f"{p.get('prev_grade') or 'not recorded'} / {migration}"],
                      ["Covenant breaches", f"{len(breaches)} of {len(cov)}"],
                      ["Early-warning signals triggered", f"{nt} of {len(ews)}"],
                      ["Review recommendation", rec]])
    doc.add_heading("2. Financial performance", 1)
    W._paras(doc, AN.income_commentary(b)[:2] + AN.balance_commentary(b)[:1])
    yrs, rows = _key_fin_rows(b)
    W._table(doc, ["Particulars (₹ Cr)"] + yrs, rows, font_size=8.5)
    pva = CT.plan_vs_actual(b, p)
    doc.add_heading("3. Projections versus actuals", 1)
    if pva:
        W._paras(doc, ["Actual results are compared with the projections accepted at the last sanction / renewal. Shortfalls above 15% "
                       "indicate that the original assessment was optimistic and the limits should be re-assessed."])
        W._table(doc, ["Item", f"Projected {years[0]}", f"Actual {years[0]}", "Variance", "Variance %"],
                 [([a, W._fmtv(pv), W._fmtv(av), W._fmtv(v), f"{vp:+.1%}"], {4: "FEE2E2" if vp < -0.15 else "DCFCE7" if vp >= 0 else "FEF3C7"})
                  for a, pv, av, v, vp in pva], widths=[4, 3.3, 3.3, 3.2, 3.2], font_size=9)
    else:
        W._paras(doc, ["Previous projections were not entered. Enter last year's projected revenue and PAT in the Proposal tab to compare."])
    doc.add_heading("4. Covenant compliance", 1)
    W._table(doc, ["Covenant", "Threshold", f"Actual {years[0]}", "Status"],
             [(c, {3: "DCFCE7" if c[3] == "Complied" else "FEE2E2"}) for c in cov], widths=[5, 4, 3, 5], font_size=9)
    doc.add_heading("5. Conduct of account", 1)
    W._table(doc, ["Parameter", "Position"],
             [["Overdues / DPD", CT.sma_status(p.get("dpd"))],
              ["Stock statements submitted on time", "No - delays reported" if p.get("stock_stmt_delay") else "Yes / not reported"],
              ["Cheque returns (12 months)", str(p.get("cheque_returns") or 0)],
              ["Ad hoc / overdrawing", "Yes" if p.get("adhoc_overdrawn") else "No / not reported"],
              ["Routing of sales through our account", "Not routed / partial" if p.get("non_routing") else "Routed / not reported"],
              ["Cash withdrawals", "High" if p.get("cash_withdrawals_high") else "Normal / not reported"]], widths=[7, 10], font_size=9)
    dp = CT.drawing_power(b, p)
    W._paras(doc, [f"Drawing power from the latest balance sheet: ₹{dp['dp']:,.2f} Cr against a limit of ₹{dp['limit']:,.2f} Cr."])
    RS.ews_section(doc, b, p, num="6")
    if b.get("comments"):
        doc.add_heading("7. Officer's observations", 1)
        W._blt(doc, [f"{c.get('topic', 'General')}: {c.get('text', '')}" for c in b["comments"]], 9.5)
    doc.add_heading("8. Recommendation", 1)
    W._paras(doc, [f"{rec}. " + ("Covenant breaches: " + "; ".join(c[0] for c in breaches) + ". " if breaches else "")
                   + (f"Early-warning signals requiring explanation: {', '.join(e[0].split(' /')[0].lower() for e in ews if e[1] == 'Triggered')}." if nt else "")])
    _signoff(doc)
    return RS.save(doc)


# ------------------------------------------------------------------ 3. early-warning & fraud-risk report
def ews_report(b, narr, officer="", branch=""):
    p = CT.proposal(b)
    fz = b["forensics"]
    doc = RS.new_doc("Early-Warning & Fraud-Risk Assessment", "RBI EWS screening, forensic accounting tests and fund-flow review", b, officer, branch)
    ews, nt = CT.rbi_ews(b, p)
    doc.add_heading("1. Conclusion", 1)
    W._paras(doc, [f"{nt} early-warning signal(s) are triggered and the forensic review raised "
                   f"{sum(1 for f in fz['flags'] if f['severity'] == 'High')} high-severity and {sum(1 for f in fz['flags'] if f['severity'] == 'Medium')} "
                   f"medium-severity observations. {CT.ews_action(nt)}"])
    RS.ews_section(doc, b, p, num="2")
    RS.screening_section(doc, b, num="2A")
    doc.add_heading("3. Forensic accounting observations", 1)
    W._table(doc, ["#", "Severity", "Observation", "Evidence", "Action required"],
             [([str(i), f["severity"], f["title"], f["evidence"][:260], f["action"]], {1: W.SEV_FILL.get(f["severity"])})
              for i, f in enumerate(_flags_sorted(b), 1)] or [["-", "-", "No observations", "-", "-"]],
             widths=[0.7, 1.6, 4, 6.2, 4.5], font_size=7.5)
    su = AN.sources_uses(b)
    if su:
        doc.add_heading("4. Fund flow and end-use", 1)
        W._paras(doc, AN.sources_uses_commentary(su))
        W._table(doc, ["Sources of funds", "₹ Cr", "Uses of funds", "₹ Cr"],
                 [[(su["sources"][i][0] if i < len(su["sources"]) else ""), (W._fmtv(su["sources"][i][1]) if i < len(su["sources"]) else ""),
                   (su["uses"][i][0] if i < len(su["uses"]) else ""), (W._fmtv(su["uses"][i][1]) if i < len(su["uses"]) else "")]
                  for i in range(min(max(len(su["sources"]), len(su["uses"])), 10))], widths=[6, 2.5, 6, 2.5], font_size=8.5)
    doc.add_heading("5. Statistical indicators", 1)
    W._keyvalue(doc, [["Beneish M-score (> -1.78 suggests possible manipulation)", f"{fz.get('beneish'):.2f}" if fz.get("beneish") is not None else "n.a."],
                      ["Piotroski F-score", f"{fz.get('piotroski')}/9" if fz.get("piotroski") is not None else "n.a."],
                      ["Altman Z'' (emerging markets)", f"{fz.get('altman_z'):.2f}" if fz.get("altman_z") is not None else "n.a."]])
    eq = CT.earnings_quality(b)
    if eq:
        W._table(doc, ["Earnings quality"] + [e["year"] for e in reversed(eq)],
                 [["Accrual ratio"] + [RS.pc(e["accrual_ratio"]) for e in reversed(eq)], ["CFO / PAT"] + [RS.xr(e["cfo_pat"]) for e in reversed(eq)]],
                 font_size=8.5)
    rps = (b.get("notes") or {}).get("related_parties") or []
    if rps:
        doc.add_heading("6. Related-party exposure", 1)
        W._table(doc, ["Party", "Relationship", "Transaction", "₹ Cr"],
                 [[x.get("name"), x.get("relationship"), x.get("transaction"), W._fmtv(x.get("amount"))] for x in rps[:25]], font_size=8)
    doc.add_heading("7. Recommended actions", 1)
    acts = ["Obtain written explanations from the borrower for each triggered signal and high-severity observation.",
            "Scrutinise 12 months' bank statements for round-tripping, related-party transfers, cash withdrawals and cheque returns.",
            "Run an MCA charge search and CRILC / CIBIL check for undisclosed borrowings.",
            "Obtain the statutory auditor's end-use certificate for borrowed funds."]
    if nt >= 4:
        acts.insert(0, "Place the account before the bank's fraud-risk / RFA committee and consider a forensic audit as required by the RBI framework.")
    W._blt(doc, acts + list(dict.fromkeys(f["action"] for f in _flags_sorted(b) if f["severity"] == "High")), 9.5)
    _signoff(doc)
    return RS.save(doc)


# ------------------------------------------------------------------ 4. query letter to borrower
def query_letter(b, narr, officer="", branch=""):
    doc = RS.new_doc("Letter to Borrower", "Clarifications and documents required", b, officer, branch, confidential=False, banner=False)
    import os as _os
    W._p(doc, ", ".join(x for x in [_os.getenv("BANK_NAME", ""), branch] if x) or "Bank name and branch: set them in ⚙ Settings",
         11, bold=True, color=W.NAVY, space_after=10)
    W._p(doc, f"Ref: CI/{dt.date.today():%Y}/{(b.get('company') or 'X')[:3].upper()}/____                                   Date: {dt.date.today():%d %B %Y}", 10)
    W._p(doc, f"To,\nThe Managing Director / Chief Financial Officer\n{b.get('company')}\n", 10.5)
    W._p(doc, "Subject: Clarifications and documents required for processing your credit proposal", 10.5, bold=True)
    W._paras(doc, ["Dear Sir / Madam,",
                   f"We refer to your proposal for credit facilities and the financial statements for {', '.join(reversed(b['years']))} submitted to us. "
                   "To complete our appraisal, we request your clarifications on the points below, supported by documents where indicated, "
                   "within 15 days of this letter."])
    doc.add_heading("A. Clarifications on the financial statements", 2)
    qs = []
    for f in _flags_sorted(b):
        internal = any(w in (f["title"] + f.get("category", "")).lower() for w in ("beneish", "altman", "manipulation", "news", "market intelligence", "distress"))
        if f["severity"] in ("High", "Medium") and not internal:
            qs.append(f"{f['title']}: our review notes that {f['evidence'][:220].rstrip('.')}. Please explain the reasons and provide supporting details.")
    an = AN.anomalies(b)
    for a in an[:6]:
        qs.append(f"{a[0]} moved from {a[1]} to {a[2]} ({a[3]}). Please provide a break-up and the reason for the change.")
    for i, q in enumerate(qs or ["No specific clarifications at this stage."], 1):
        p_ = doc.add_paragraph()
        r_ = p_.add_run(f"{i}. {q}")
        r_.font.size = Pt(10)
    doc.add_heading("B. Documents required", 2)
    docs = ["Audited financial statements with all schedules, notes and the tax audit report (Form 3CD) for the last three years.",
            "Provisional financials / management accounts for the current year to date.",
            "Projected balance sheet and profit & loss for the loan tenor with assumptions (CMA data).",
            "GST returns (GSTR-1 and GSTR-3B) for the last 12 months.",
            "Statements of all bank accounts for the last 12 months.",
            "Age-wise list of debtors and creditors with the top-10 parties; confirmations for related-party balances.",
            "Details of all existing borrowings with sanction letters and repayment schedules; latest stock statement.",
            "Shareholding pattern and details of pledged shares; net-worth statements of guarantors.",
            "Valuation and title reports for the security offered."]
    for i, d in enumerate(docs, 1):
        p_ = doc.add_paragraph()
        r_ = p_.add_run(f"{i}. {d}")
        r_.font.size = Pt(10)
    W._paras(doc, ["", "We will be glad to discuss any of these points at a meeting at your convenience.", "Yours faithfully,", "",
                   f"{officer or '(Authorised signatory)'}\n{branch or ''}"])
    return RS.save(doc)


# ------------------------------------------------------------------ 5. indicative term sheet
def term_sheet(b, narr, officer="", branch=""):
    p = CT.proposal(b)
    doc = RS.new_doc("Indicative Term Sheet", "Draft - subject to sanction by the competent authority and execution of documents", b, officer, branch)
    fac = []
    if p.get("tl_amount"):
        _, inst, n = CT.amortisation(p["tl_amount"], p.get("tl_tenor_years"), p.get("tl_moratorium_months") or 0, p.get("tl_rate") or 0.1)
        fac.append(["Term loan", f"₹{p['tl_amount']:,.2f} Cr", f"{p.get('tl_tenor_years')} years incl. {p.get('tl_moratorium_months')} m moratorium",
                    f"{n} quarterly instalments of ₹{inst:,.3f} Cr", RS.pc(p.get("tl_rate"), 2)])
    if p.get("cc_amount"):
        fac.append(["Cash credit (working capital)", f"₹{p['cc_amount']:,.2f} Cr", "12 months, renewable", "On demand; drawing power basis", RS.pc(p.get("cc_rate"), 2)])
    if p.get("nfb_amount"):
        fac.append(["Bank guarantee / LC (non-fund)", f"₹{p['nfb_amount']:,.2f} Cr", "12 months, renewable", "-", "Commission as per schedule"])
    doc.add_heading("1. Facilities", 1)
    W._table(doc, ["Facility", "Amount", "Tenor", "Repayment", "Rate (p.a.)"], fac or [[p.get("facility") or "-", "-", "-", "-", "-"]],
             widths=[4, 2.5, 3.5, 4.5, 2.5], font_size=8.5)
    rp = CT.risk_pricing(b, p)
    W._note(doc, f"Rates are linked to the bank's benchmark ({RS.pc(rp['base_rate'], 2)}) and reset with it"
                 + (f"; risk-based minimum {RS.pc(rp['floor'], 2)}." if rp.get("configured") else "; risk-based minimum not computed (PD scale not configured)."))
    doc.add_heading("2. Purpose", 1)
    W._paras(doc, [p.get("purpose") or "To be specified."])
    doc.add_heading("3. Security", 1)
    W._table(doc, ["Type", "Details"], [["Primary", p.get("primary_security")],
                                        ["Collateral", f"{p.get('collateral_type')} - market value ₹{p.get('collateral_value') or 0:,.2f} Cr"],
                                        ["Guarantee", p.get("guarantee")]], widths=[3.5, 13.5], font_size=9)
    doc.add_heading("4. Margins", 1)
    W._table(doc, ["Item", "Margin"], [["Stock (paid)", RS.pc(p.get("stock_margin"), 0)], ["Receivables up to 90 days", RS.pc(p.get("debtor_margin"), 0)],
                                       ["Term loan (promoter contribution)", "25% of project cost (to be confirmed)"]], widths=[8, 9], font_size=9)
    doc.add_heading("5. Financial covenants", 1)
    W._table(doc, ["Covenant", "Threshold", "Test"], [[c[0], c[1], "Annually on audited accounts"] for c in AN.covenants(b)], widths=[6, 5, 6], font_size=9)
    doc.add_heading("6. Conditions precedent to first disbursement", 1)
    W._blt(doc, ["Acceptance of sanction letter and execution of loan and security documents.",
                 "Creation of charge on primary and collateral security; filing of CHG-1 with ROC within 30 days.",
                 "Valuation and legal title reports satisfactory to the bank; insurance of assets with bank clause.",
                 "Board resolution authorising the borrowing; KYC of company, directors and guarantors.",
                 "Satisfactory CIBIL / CRILC reports; no-objection from existing lenders for pari-passu charge where applicable.",
                 "Promoter contribution brought in upfront and certified by a chartered accountant."], 9.5)
    doc.add_heading("7. Conditions subsequent and ongoing", 1)
    W._blt(doc, (narr.get("suggested_covenants") if narr.get("_source") == "AI" else None) or [
        "Monthly stock and book-debt statements within 20 days of month-end; drawing power on paid stock and receivables up to 90 days.",
        "Annual stock audit and unit inspection; quarterly financial information (QIS / FFR).",
        "All sales proceeds routed through the bank's accounts; no diversion to group or related entities without consent.",
        "End-use certificate from the statutory auditor within 60 days of each term-loan disbursement.",
        "No change in management control or shareholding without prior consent; immediate intimation of rating action or litigation."], 9.5)
    doc.add_heading("8. Other terms", 1)
    W._blt(doc, ["Prepayment: as per the bank's policy.", "Default interest: 2% p.a. over the applicable rate on overdue amounts.",
                 "Validity: this indicative term sheet is valid for 30 days and is not a commitment to lend."], 9.5)
    return RS.save(doc)


# ------------------------------------------------------------------ 6. covenant compliance report
def covenant_report(b, narr, officer="", branch=""):
    p = CT.proposal(b)
    doc = RS.new_doc("Covenant Compliance Report", "Financial covenants tested on audited accounts, with projected headroom", b, officer, branch)
    cov = AN.covenants(b)
    breaches = [c for c in cov if c[3].startswith("Currently breached")]
    doc.add_heading("1. Summary", 1)
    W._paras(doc, [f"{len(cov) - len(breaches)} of {len(cov)} financial covenants are complied with on the {b['years'][0]} accounts."
                   + (" Breached: " + "; ".join(f"{c[0]} ({c[2]} vs {c[1].lower()})" for c in breaches) + ". A cure plan, waiver or "
                      "re-pricing should be considered." if breaches else "")])
    doc.add_heading("2. Tested covenants", 1)
    W._table(doc, ["Covenant", "Threshold", f"Actual {b['years'][0]}", "Status"],
             [(c, {3: "DCFCE7" if c[3] == "Complied" else "FEE2E2"}) for c in cov], widths=[5, 4, 3, 5], font_size=9)
    rows, _ = CT.projections(b, p)
    doc.add_heading("3. Projected covenant headroom", 1)
    W._paras(doc, ["Base-case projections show whether the covenants are likely to remain in compliance over the facility tenor."])
    lim = {"DSCR (min 1.25x)": ("dscr", 1.25, "min"), "TOL / TNW (max 3.0x)": ("tol_tnw", 3.0, "max"),
           "Debt / EBITDA (max 3.5x)": ("debt_ebitda", 3.5, "max"), "Interest coverage (min 2.0x)": ("icr", 2.0, "min")}
    tab = []
    for lab, (k, t, op) in lim.items():
        cells, fills = [lab], {}
        for i, r_ in enumerate(rows, 1):
            v = r_[k]
            cells.append(RS.xr(v))
            if v is not None:
                fills[i] = "DCFCE7" if (v >= t if op == "min" else v <= t) else "FEE2E2"
        tab.append((cells, fills))
    W._table(doc, ["Covenant"] + [r_["year"] for r_ in rows], tab, font_size=8.5)
    doc.add_heading("4. Information and conduct covenants", 1)
    W._table(doc, ["Covenant", "Status", "Remarks"],
             [["Monthly stock / book-debt statement", "Delayed" if p.get("stock_stmt_delay") else "☐ Verify", ""],
              ["Quarterly financial information (QIS / FFR)", "☐ Verify", ""], ["Annual stock audit", "☐ Verify", ""],
              ["Insurance of assets with bank clause", "☐ Verify", ""], ["Routing of sales through the bank", "Not complied" if p.get("non_routing") else "☐ Verify", ""],
              ["No related-party lending without consent", "☐ Verify", ""], ["End-use certificate", "☐ Verify", ""]], widths=[7, 3.5, 6.5], font_size=9)
    _signoff(doc)
    return RS.save(doc)


# ------------------------------------------------------------------ 7. ESG & financed emissions note
def esg_note(b, narr, officer="", branch=""):
    doc = RS.new_doc("ESG & Financed Emissions Note", "BRSR review and PCAF financed-emissions calculation", b, officer, branch)
    br = b.get("brsr") or {}
    pcf = b.get("pcaf") or {}
    y0 = b["years"][0]
    rev = g(b["data"][y0], "revenue")
    doc.add_heading("1. Summary", 1)
    s12 = (br.get("scope1_tco2e") or 0) + (br.get("scope2_tco2e") or 0)
    W._paras(doc, [br.get("summary") or "BRSR not provided - ESG indicators are not assessed.",
                   (f"Emission intensity: {s12 / rev:,.1f} tCO2e per ₹ crore of revenue (Scope 1+2)." if s12 and rev else "")])
    if br:
        doc.add_heading("2. BRSR indicators", 1)
        items = [("Scope 1 (tCO2e)", "scope1_tco2e"), ("Scope 2 (tCO2e)", "scope2_tco2e"), ("Scope 3 (tCO2e)", "scope3_tco2e"),
                 ("Energy consumed (GJ)", "energy_total_gj"), ("Renewable energy share (%)", "energy_renewable_pct"),
                 ("Water withdrawal (kL)", "water_withdrawal_kl"), ("Waste generated (t)", "waste_generated_t"),
                 ("GHG assurance", "ghg_assurance"), ("Net-zero target", "net_zero_target"), ("LTIFR", "ltifr_employees"), ("Fatalities", "fatalities")]
        W._table(doc, ["Indicator", "Reported"], [[lab, f"{br.get(k):,.2f}" if isinstance(br.get(k), (int, float)) else str(br.get(k))]
                                                  for lab, k in items if br.get(k) not in (None, "")] or [["-", "-"]], widths=[8, 8], font_size=9)
        if br.get("esg_red_flags"):
            W._blt(doc, [f"[{x.get('severity')}] {x.get('title')}: {x.get('evidence')}" for x in br["esg_red_flags"]], 9.5)
    doc.add_heading("3. PCAF financed emissions", 1)
    if pcf and pcf.get("attribution_factor") is not None:
        W._table(doc, ["Item", "Value"], [["Asset class", (pcf.get("inputs") or {}).get("asset_class", "-")], ["Outstanding (₹ Cr)", W._fmtv((pcf.get("inputs") or {}).get("outstanding_cr"))],
                                          ["Denominator", str(pcf.get("denominator_name"))], ["Attribution factor", f"{pcf['attribution_factor']:.4%}"],
                                          ["Company Scope 1+2 (tCO2e)", f"{pcf.get('company_scope12') or 0:,.0f}"],
                                          ["Financed Scope 1+2 (tCO2e)", f"{pcf.get('financed_scope12') or 0:,.0f}"],
                                          ["Financed Scope 3 (tCO2e)", f"{pcf.get('financed_scope3') or 0:,.0f}" if pcf.get("financed_scope3") else "-"],
                                          ["PCAF data-quality score", str(pcf.get("data_quality_score"))]], widths=[8, 8], font_size=9)
        if pcf.get("warning"):
            W._note(doc, pcf["warning"])
    else:
        W._paras(doc, ["PCAF inputs not provided - enter exposure and emissions (or fetch them from the web) in the BRSR & PCAF tab."])
    doc.add_heading("4. Transition-risk considerations", 1)
    W._blt(doc, ["Carbon-intensive borrowers face rising costs from carbon pricing, the Indian Carbon Credit Trading Scheme and buyer requirements (e.g. EU CBAM for exports).",
                 "Seek a decarbonisation plan with targets and capex where emission intensity is high or emissions are undisclosed.",
                 "Improve the PCAF data-quality score over time by obtaining verified, borrower-reported emissions."], 9.5)
    return RS.save(doc)


REPORT_TYPES = {
    "Credit Appraisal Memorandum (full)": (None, "The complete 20-page memorandum: analysis, projections, forensic, ESG, pricing, covenants."),
    "Credit Committee Note": (committee_note, "3-page summary for the sanctioning committee with sign-off block."),
    "Annual Review / Renewal Note": (annual_review, "Performance vs last year and projections, covenants, account conduct, EWS, renewal call."),
    "Early-Warning & Fraud-Risk Report": (ews_report, "RBI EWS screening, forensic findings, fund flow, statistical scores, actions."),
    "Query Letter to Borrower": (query_letter, "Ready-to-send letter listing clarifications and documents required."),
    "Indicative Term Sheet": (term_sheet, "Facilities, pricing, security, margins, covenants, conditions precedent / subsequent."),
    "Covenant Compliance Report": (covenant_report, "Covenants tested on actuals plus projected headroom year by year."),
    "ESG & Financed Emissions Note": (esg_note, "BRSR indicators, emission intensity, PCAF calculation, transition risk."),
}


def build(name, bundle, narr, officer="", branch=""):
    from . import checker
    narr, bundle["checks"] = checker.run(bundle, narr)  # every report type goes through the same logic checker
    fn = REPORT_TYPES[name][0]
    if fn is None:
        return W.build(bundle, narr, officer=officer, branch=branch)
    return fn(bundle, narr, officer=officer, branch=branch)
