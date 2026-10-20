"""Extra workbook tabs (added after the three core tabs):
 4  Projections & DSCR   - live formulas driven by yellow assumption cells (change growth / margin / rate and DSCR recalculates)
 5  Repayment Schedule   - quarterly term-loan schedule with moratorium (formulas)
 6  Scenarios            - base / upside / downside / severe stress, break-even EBITDA fall
 7  WC Assessment & DP   - Tandon I & II, turnover method, drawing power
 8  Security & Pricing   - security cover, PD / LGD / EL, hurdle rate, RAROC
 9  RBI EWS & Regulatory - early-warning signals and exposure / IRACP checks
 10 CMA Analysis         - CMA-format Forms II, III, IV and VI from the audited spread
"""
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from . import analyst as AN
from . import credit_tools as CT
from . import ratios as RATIOMOD
from . import report_excel as X
from .schema import g

PCT = '0.0%;(0.0%);"-"'
XF = '0.00"x";(0.00"x");"-"'


def _title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(name=X.FONT, bold=True, size=13, color=X.NAVY)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = X.F_NOTE


def _put(ws, r, c, v, fmt=None, font=None, fill=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.border = X.BOX
    cell.font = font or X.F_CALC
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    return cell


def _rows(ws, start, header, rows, widths, fmts=None):
    X._header(ws, start, header, widths)
    r = start + 1
    for row in rows:
        for j, v in enumerate(row, 1):
            f = (fmts or {}).get(j) if isinstance(v, (int, float)) else None
            _put(ws, r, j, v, f or (X.NUM if isinstance(v, float) else None))
        r += 1
    return r


def _traffic(ws, rng, good, bad, higher_better=True):
    if higher_better:
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=[str(good)], fill=X.FILL_GOOD))
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=[str(bad)], fill=X.FILL_BAD))
    else:
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThanOrEqual", formula=[str(good)], fill=X.FILL_GOOD))
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=[str(bad)], fill=X.FILL_BAD))


def add_tabs(wb, b):
    p = CT.proposal(b)
    y0 = b["years"][0]
    c = b["data"][y0]
    miss_ = CT.core_missing(b)

    rows_py, a = CT.projections(b, p) if not miss_ else ([], {"status": "NOT COMPUTED", "reason": "spread is missing: " + ", ".join(miss_)})
    if miss_ or a.get("status") == "NOT COMPUTED":
        for nm_ in ("Projections & DSCR", "Repayment Schedule", "Scenarios", "WC Assessment & DP"):
            if nm_ == "WC Assessment & DP" and not miss_:
                continue  # working capital does not depend on the growth rate's source; built below
            ws_ = wb.create_sheet(nm_)
            _title(ws_, f"{nm_} - NOT COMPUTED - INPUT REQUIRED", f"{a.get('reason')}. These sheets would otherwise rest on blanks treated as "
                   "zero or on a growth rate nobody has verified. Enter the missing input in the app and regenerate the workbook.")
    if not miss_ and a.get("status") != "NOT COMPUTED":
        n = len(rows_py)
        # ------------------------------------------------------------ 5. repayment schedule (built first, projections refer to it)
        rs = wb.create_sheet("Repayment Schedule")
        # ------------------------------------------------------------ 4. projections
        ws = wb.create_sheet("Projections & DSCR", index=wb.sheetnames.index("Repayment Schedule"))
        _title(ws, f"{b.get('company')} - projected financials and DSCR (₹ Cr)",
               "Yellow cells are assumptions: change them and every projection, DSCR and coverage figure recalculates.")
        inputs = [
            ("Base revenue (latest audited)", g(c, "revenue"), X.NUM), ("Revenue growth p.a.", a["growth"], PCT),
            ("EBITDA margin", a["margin"], PCT), ("Depreciation (existing assets)", g(c, "depreciation"), X.NUM),
            ("Other income p.a.", g(c, "other_income"), X.NUM), ("Existing term debt", a["existing_ltd"], X.NUM),
            ("Existing term debt - remaining years", a["existing_ltd_years"], "0"), ("Interest rate on existing debt", a["debt_rate"], PCT),
            ("Working-capital borrowings", a["wc_debt"], X.NUM), ("Interest rate on working capital", a["wc_rate"], PCT),
            ("New term loan", p.get("tl_amount") or 0.0, X.NUM), ("New term loan rate", (p.get("tl_rate") or 0.1) + (p.get("rate_shift") or 0.0), PCT),
            ("Tax rate", CT.TAX, PCT), ("Depreciation on new assets", a["new_dep"], X.NUM), ("Opening net worth", g(c, "total_equity"), X.NUM),
            ("Term-loan tenor (years)", p.get("tl_tenor_years") or 0, "0"), ("Moratorium (months)", p.get("tl_moratorium_months") or 0, "0"),
        ]
        X._header(ws, 3, ["Assumption", "Value", "Source (from the assumption register)"], [44, 14, 70])
        _reg = {r_[0]: r_[2] for r_ in a.get("register") or []}
        _src_of = {"Base revenue (latest audited)": "Base revenue", "Revenue growth p.a.": "Revenue growth (a year)", "EBITDA margin": "EBITDA margin",
                   "Depreciation (existing assets)": "Depreciation", "Other income p.a.": "Other income",
                   "Existing term debt": "Existing term debt / remaining tenor", "Existing term debt - remaining years": "Existing term debt / remaining tenor",
                   "Interest rate on existing debt": "Interest rate on existing debt", "Working-capital borrowings": "Working-capital borrowing / rate",
                   "Interest rate on working capital": "Working-capital borrowing / rate", "New term loan": "New term loan amount",
                   "New term loan rate": "New term loan rate", "Tax rate": "Tax rate", "Term-loan tenor (years)": "Tenor / moratorium",
                   "Moratorium (months)": "Tenor / moratorium"}
        A = {}
        for i, (lab, v, f) in enumerate(inputs):
            r = 4 + i
            _put(ws, r, 1, lab)
            _put(ws, r, 2, v, f, X.F_IN, X.FILL_INPUT)
            src_ = _reg.get(_src_of.get(lab, ""), "verified spread" if lab in ("Opening net worth",) else "model assumption")
            _put(ws, r, 3, src_, font=X.F_NOTE, fill=X.FILL_WATCH if ("approximation" in src_ or "not reviewed" in src_ or "NOT DISCLOSED" in src_) else None)
            A[lab] = f"$B${r}"
        ref = lambda lab: A[lab]
        top = 4 + len(inputs) + 1
        X._header(ws, top, ["Projection"] + [f"Y+{i}" for i in range(1, n + 1)], [44] + [14] * n)
        labels = ["Revenue", "EBITDA", "Other income", "Depreciation", "Existing TL - opening", "Existing TL - repayment", "Existing TL - closing",
                  "Existing TL - interest", "New TL - interest", "New TL - repayment", "Working-capital interest", "Total interest",
                  "Profit before tax", "Tax", "Profit after tax", "Cash accruals (PAT + dep.)", "Term-debt interest", "Term-debt principal",
                  "DSCR", "Interest coverage", "Surplus after term repayment", "Net worth (closing)"]
        R_ = {lab: top + 1 + i for i, lab in enumerate(labels)}
        for lab, r in R_.items():
            _put(ws, r, 1, lab, font=X.F_BOLD if lab in ("DSCR", "Profit after tax", "Total interest") else X.F_CALC)
        for j in range(n):
            col = get_column_letter(2 + j)
            prev = get_column_letter(1 + j)
            yr = j + 1
            f = {
                "Revenue": f"={ref('Base revenue (latest audited)')}*(1+{ref('Revenue growth p.a.')})" if j == 0 else f"={prev}{R_['Revenue']}*(1+{ref('Revenue growth p.a.')})",
                "EBITDA": f"={col}{R_['Revenue']}*{ref('EBITDA margin')}",
                "Other income": f"={ref('Other income p.a.')}",
                "Depreciation": f"={ref('Depreciation (existing assets)')}+{ref('Depreciation on new assets')}",
                "Existing TL - opening": f"={ref('Existing term debt')}" if j == 0 else f"={prev}{R_['Existing TL - closing']}",
                "Existing TL - repayment": f"=MIN({ref('Existing term debt')}/MAX({ref('Existing term debt - remaining years')},1),{col}{R_['Existing TL - opening']})",
                "Existing TL - closing": f"={col}{R_['Existing TL - opening']}-{col}{R_['Existing TL - repayment']}",
                "Existing TL - interest": f"=({col}{R_['Existing TL - opening']}-{col}{R_['Existing TL - repayment']}/2)*{ref('Interest rate on existing debt')}",
                "New TL - interest": f"=SUMIF('Repayment Schedule'!$B:$B,{yr},'Repayment Schedule'!$E:$E)",
                "New TL - repayment": f"=SUMIF('Repayment Schedule'!$B:$B,{yr},'Repayment Schedule'!$F:$F)",
                "Working-capital interest": f"={ref('Working-capital borrowings')}*{ref('Interest rate on working capital')}",
                "Total interest": f"={col}{R_['Existing TL - interest']}+{col}{R_['New TL - interest']}+{col}{R_['Working-capital interest']}",
                "Profit before tax": f"={col}{R_['EBITDA']}+{col}{R_['Other income']}-{col}{R_['Depreciation']}-{col}{R_['Total interest']}",
                "Tax": f"=MAX({col}{R_['Profit before tax']},0)*{ref('Tax rate')}",
                "Profit after tax": f"={col}{R_['Profit before tax']}-{col}{R_['Tax']}",
                "Cash accruals (PAT + dep.)": f"={col}{R_['Profit after tax']}+{col}{R_['Depreciation']}",
                "Term-debt interest": f"={col}{R_['Existing TL - interest']}+{col}{R_['New TL - interest']}",
                "Term-debt principal": f"={col}{R_['Existing TL - repayment']}+{col}{R_['New TL - repayment']}",
                "DSCR": f"=IFERROR(({col}{R_['Profit after tax']}+{col}{R_['Depreciation']}+{col}{R_['Term-debt interest']})/({col}{R_['Term-debt interest']}+{col}{R_['Term-debt principal']}),\"-\")",
                "Interest coverage": f"=IFERROR(({col}{R_['EBITDA']}+{col}{R_['Other income']}-{col}{R_['Depreciation']})/{col}{R_['Total interest']},\"-\")",
                "Surplus after term repayment": f"={col}{R_['Cash accruals (PAT + dep.)']}-{col}{R_['Term-debt principal']}",
                "Net worth (closing)": (f"={ref('Opening net worth')}+{col}{R_['Profit after tax']}" if j == 0 else f"={prev}{R_['Net worth (closing)']}+{col}{R_['Profit after tax']}"),
            }
            for lab, formula in f.items():
                _put(ws, R_[lab], 2 + j, formula, XF if lab in ("DSCR", "Interest coverage") else X.NUM)
        last = get_column_letter(1 + n)
        _traffic(ws, f"B{R_['DSCR']}:{last}{R_['DSCR']}", 1.25, 1.10)
        sr = R_["Net worth (closing)"] + 2
        _put(ws, sr, 1, "Average DSCR (total cash available / total debt service)", font=X.F_BOLD)
        _put(ws, sr, 2, f"=IFERROR((SUM(B{R_['Profit after tax']}:{last}{R_['Profit after tax']})+SUM(B{R_['Depreciation']}:{last}{R_['Depreciation']})"
                        f"+SUM(B{R_['Term-debt interest']}:{last}{R_['Term-debt interest']}))/(SUM(B{R_['Term-debt interest']}:{last}{R_['Term-debt interest']})"
                        f"+SUM(B{R_['Term-debt principal']}:{last}{R_['Term-debt principal']})),\"-\")", XF, X.F_BOLD)
        _put(ws, sr + 1, 1, "Minimum DSCR", font=X.F_BOLD)
        _put(ws, sr + 1, 2, f"=MIN(B{R_['DSCR']}:{last}{R_['DSCR']})", XF, X.F_BOLD)
        _traffic(ws, f"B{sr}:B{sr}", 1.25, 1.10)
        _traffic(ws, f"B{sr + 1}:B{sr + 1}", 1.10, 1.0)
        ws.cell(row=sr + 3, column=1, value="Bank norms: average DSCR >= 1.25x, minimum >= 1.10x (check the bank's policy).").font = X.F_NOTE
        ws.freeze_panes = f"B{top + 1}"

        # repayment schedule formulas
        _title(rs, "New term loan - quarterly repayment schedule (₹ Cr)", "Amount, rate, tenor and moratorium link to the assumptions on 'Projections & DSCR'.")
        P = "'Projections & DSCR'!"
        X._header(rs, 3, ["Quarter", "Year", "Opening", "Instalment due", "Interest", "Principal", "Closing"], [10, 8, 14, 14, 14, 14, 14])
        q_total = int(round((p.get("tl_tenor_years") or 0) * 4))
        amt, ten, mor, rate = (P + ref("New term loan"), P + ref("Term-loan tenor (years)"), P + ref("Moratorium (months)"), P + ref("New term loan rate"))
        rs["I3"], rs["J3"] = "Instalment", f"=IFERROR({amt}/({ten}*4-ROUND({mor}/3,0)),0)"
        rs["I3"].font = X.F_BOLD
        rs["J3"].number_format = X.NUM
        for q in range(1, max(q_total, 1) + 1):
            r = 3 + q
            _put(rs, r, 1, q, "0")
            _put(rs, r, 2, f"=ROUNDUP(A{r}/4,0)", "0")
            _put(rs, r, 3, f"={amt}" if q == 1 else f"=G{r - 1}", X.NUM)
            _put(rs, r, 4, f"=IF(A{r}>ROUND({mor}/3,0),$J$3,0)", X.NUM)
            _put(rs, r, 5, f"=C{r}*{rate}/4", X.NUM)
            _put(rs, r, 6, f"=MIN(D{r},C{r})", X.NUM)
            _put(rs, r, 7, f"=C{r}-F{r}", X.NUM)
        if not p.get("tl_amount"):
            rs.cell(row=6 + max(q_total, 1), column=1, value="No new term loan entered - enter it in the Proposal tab of the app.").font = X.F_NOTE

        # ------------------------------------------------------------ 6. scenarios
        sc = wb.create_sheet("Scenarios")
        _title(sc, "Scenario analysis over the projection period", "Each scenario flexes growth, EBITDA margin and interest rates together.")
        srows = [[s["scenario"], s["growth"], s["margin_shift"], s["rate_shift"], s["avg_dscr"], s["min_dscr"], s["y1_icr"], s["cum_surplus"]]
                 for s in CT.scenarios(b, p)]
        end = _rows(sc, 3, ["Scenario", "Growth", "Margin shift", "Rate shift", "Avg DSCR", "Min DSCR", "Y+1 ICR", "Cum. surplus ₹ Cr"], srows,
                    [18, 11, 13, 11, 11, 11, 11, 16], {2: PCT, 3: PCT, 4: PCT, 5: XF, 6: XF, 7: XF, 8: X.NUM})
        _traffic(sc, f"F4:F{end - 1}", 1.10, 1.0)
        be = CT.breakeven(b, p)
        sc.cell(row=end + 1, column=1, value="Break-even: maximum uniform fall in projected EBITDA with minimum DSCR >= 1.00x").font = X.F_BOLD
        _put(sc, end + 1, 6, be if be is not None else "n.a.", PCT, X.F_BOLD)

    if not miss_:
        # ------------------------------------------------------------ 7. WC assessment & DP
        wcs = wb.create_sheet("WC Assessment & DP")
        w = CT.wc_methods(b, p)
        _title(wcs, "Working-capital assessment (₹ Cr)", f"Projected turnover ₹{w['proj_turnover']:,.2f} Cr; growth {w['growth']:.1%} ({w.get('growth_note')}); "
                   + ("turnover method applies (limits up to ₹5 Cr)." if w.get("nayak_applies") else "turnover method shown for reference only (applies to MSE limits up to ₹5 Cr)."))
        end = _rows(wcs, 3, ["Particulars", "Tandon Method I", "Tandon Method II", "Turnover method"],
                    [[r_[0]] + [v if v is not None else "-" for v in r_[1:]] for r_ in w["rows"]], [52, 16, 16, 16], {2: X.NUM, 3: X.NUM, 4: X.NUM})
        for lab, v in [("Requested / existing WC limit", w["requested"]), ("Assessed eligibility", w["assessed"]), ("Headroom / (excess)", w["headroom"])]:
            _put(wcs, end, 1, lab, font=X.F_BOLD)
            _put(wcs, end, 3, v, X.NUM, X.F_BOLD)
            end += 1
        dp = CT.drawing_power(b, p)
        end = _rows(wcs, end + 2, ["Drawing power", "₹ Cr"], [[a_, v] for a_, v in dp["rows"]], [52, 16], {2: X.NUM})

    # ------------------------------------------------------------ 8. security & pricing
    sp = wb.create_sheet("Security & Pricing")
    scv = CT.security_cover(b, p)
    _title(sp, "Security cover and risk-based pricing", "PD scale, haircuts, cost of funds and capital ratio are illustrative - replace with the bank's policy.")
    end = _rows(sp, 3, ["Security", "₹ Cr"], [[a_, v] for a_, v in scv["rows"]] +
                [["Book-asset coverage (x) - indicative; charge not verified", scv["cover"] if scv["cover"] is not None else "n.a."],
                 ["Fixed-asset coverage ratio (x)", scv["facr"] if scv["facr"] is not None else "n.a."]], [60, 30], {2: X.NUM})
    sp.cell(row=end - 2, column=2).number_format = XF
    sp.cell(row=end - 1, column=2).number_format = XF
    _rows(sp, end + 1, ["Pricing component", "Value"], CT.risk_pricing(b, p)["rows"], [60, 30])

    # ------------------------------------------------------------ 9. EWS & regulatory
    ew = wb.create_sheet("RBI EWS & Regulatory")
    ews, nt = CT.rbi_ews(b, p)
    _title(ew, f"RBI early-warning signals - {nt} triggered", CT.ews_action(nt))
    end = _rows(ew, 3, ["Early-warning signal", "Status", "Basis"], ews, [70, 18, 70])
    ew.conditional_formatting.add(f"B4:B{end}", CellIsRule(operator="equal", formula=['"Triggered"'], fill=X.FILL_BAD))
    ew.conditional_formatting.add(f"B4:B{end}", CellIsRule(operator="equal", formula=['"No indication"'], fill=X.FILL_GOOD))
    ew.conditional_formatting.add(f"B4:B{end}", CellIsRule(operator="equal", formula=['"Not assessed"'], fill=X.FILL_WATCH))
    ew.conditional_formatting.add(f"B4:B{end}", CellIsRule(operator="equal", formula=['"Bank data needed"'], fill=X.FILL_WATCH))
    _rows(ew, end + 2, ["Regulatory check", "Position", "Result"], CT.regulatory(b, p), [70, 18, 70])

    # ------------------------------------------------------------ 10. CMA analysis
    cm = wb.create_sheet("CMA Analysis")
    yrs = list(reversed(b["years"]))
    d, r = b["data"], b["ratios"]
    _title(cm, "CMA-format analysis (audited years)", "Forms II (operating statement), III (balance sheet analysis), IV (holding levels) and VI (fund flow).")
    hdr = ["Particulars (₹ Cr)"] + yrs
    wid = [52] + [15] * len(yrs)

    class _Missing(Exception):
        pass

    def G(row, k):
        v = row.get(k)
        if v in (None, "") and (k in RATIOMOD.BLOCKING or k in RATIOMOD.DERIVED_KEYS):
            raise _Missing(k)
        return v or 0.0

    def _cell(fn, y):
        try:
            return fn(y)
        except _Missing:
            return "n.a. - input missing"
        except ZeroDivisionError:
            return "-"

    def line(lab, fn, fmt=X.NUM):
        return [lab] + [_cell(fn, y) for y in yrs], fmt

    form2 = [line("Net sales (revenue from operations)", lambda y: G(d[y], "revenue")),
             line("% rise / (fall) in net sales", lambda y: r[y].get("rev_growth") if r[y].get("rev_growth") is not None else "-", PCT),
             line("Cost of sales (materials + traded goods + change in stock)", lambda y: G(d[y], "cogs")),
             line("Employee cost", lambda y: G(d[y], "employee_cost")), line("Other operating expenses", lambda y: G(d[y], "other_expenses")),
             line("Operating profit before interest & depreciation (EBITDA)", lambda y: G(d[y], "ebitda")),
             line("Depreciation", lambda y: G(d[y], "depreciation")), line("Interest", lambda y: G(d[y], "finance_cost")),
             line("Other income", lambda y: G(d[y], "other_income")), line("Exceptional items", lambda y: G(d[y], "exceptional_items")),
             line("Profit before tax", lambda y: G(d[y], "pbt")), line("Provision for tax", lambda y: G(d[y], "tax_expense")),
             line("Net profit", lambda y: G(d[y], "pat")), line("PBT / net sales", lambda y: G(d[y], "pbt") / G(d[y], "revenue") if G(d[y], "revenue") else "-", PCT)]
    form3 = [line("Short-term bank borrowings for working capital", lambda y: G(d[y], "short_term_borrowings") - G(d[y], "current_maturities_ltd")),
             line("Instalments of term loans due within one year", lambda y: G(d[y], "current_maturities_ltd")),
             line("Trade payables", lambda y: G(d[y], "trade_payables")),
             line("Other current liabilities & provisions", lambda y: G(d[y], "total_current_liabilities") - G(d[y], "short_term_borrowings") - G(d[y], "trade_payables")),
             line("Total current liabilities", lambda y: G(d[y], "total_current_liabilities")),
             line("Term liabilities (long-term borrowings + leases)", lambda y: G(d[y], "long_term_borrowings") + G(d[y], "lease_liabilities_nc")),
             line("Other non-current liabilities & provisions", lambda y: G(d[y], "tol") - G(d[y], "total_current_liabilities") - G(d[y], "long_term_borrowings") - G(d[y], "lease_liabilities_nc")),
             line("Total outside liabilities", lambda y: G(d[y], "tol")), line("Net worth", lambda y: G(d[y], "total_equity")),
             line("Total current assets", lambda y: G(d[y], "total_current_assets")), line("Net fixed assets (incl. CWIP)", lambda y: G(d[y], "ppe") + G(d[y], "cwip") + G(d[y], "right_of_use")),
             line("Intangibles", lambda y: G(d[y], "goodwill") + G(d[y], "intangibles")),
             line("Other non-current assets & investments", lambda y: G(d[y], "total_assets") - G(d[y], "total_current_assets") - G(d[y], "ppe") - G(d[y], "cwip") - G(d[y], "right_of_use") - G(d[y], "goodwill") - G(d[y], "intangibles")),
             line("Total assets", lambda y: G(d[y], "total_assets")), line("Tangible net worth", lambda y: G(d[y], "tnw")),
             line("Net working capital", lambda y: G(d[y], "nwc")), line("Current ratio", lambda y: r[y].get("current_ratio") or "-", XF),
             line("TOL / TNW", lambda y: r[y].get("tol_tnw") or "-", XF)]
    form4 = [line("Raw material + stock holding (days of cost of sales)", lambda y: r[y].get("inventory_days") or "-", '#,##0" d"'),
             line("Receivables (days of sales)", lambda y: r[y].get("debtor_days") or "-", '#,##0" d"'),
             line("Trade payables (days of purchases)", lambda y: r[y].get("creditor_days") or "-", '#,##0" d"'),
             line("Inventory", lambda y: G(d[y], "inventories")), line("Receivables", lambda y: G(d[y], "trade_receivables")),
             line("Other current assets (incl. cash)", lambda y: G(d[y], "total_current_assets") - G(d[y], "inventories") - G(d[y], "trade_receivables"))]
    rr = 3
    for title, block in [("FORM II - Operating statement", form2), ("FORM III - Analysis of balance sheet", form3),
                         ("FORM IV - Current assets and holding levels", form4)]:
        cm.cell(row=rr, column=1, value=title).font = X.F_SEC
        X._header(cm, rr + 1, hdr, wid)
        rr += 2
        for vals, fmt in block:
            for j, v in enumerate(vals, 1):
                _put(cm, rr, j, v, fmt if j > 1 and isinstance(v, (int, float)) else None)
            rr += 1
        rr += 1
    su = AN.sources_uses(b)
    if su:
        cm.cell(row=rr, column=1, value=f"FORM VI - Fund flow ({su['from']} to {su['to']})").font = X.F_SEC
        rr = _rows(cm, rr + 1, ["Sources", "₹ Cr", "Uses", "₹ Cr"],
                   [[(su["sources"][i][0] if i < len(su["sources"]) else ""), (su["sources"][i][1] if i < len(su["sources"]) else ""),
                     (su["uses"][i][0] if i < len(su["uses"]) else ""), (su["uses"][i][1] if i < len(su["uses"]) else "")]
                    for i in range(max(len(su["sources"]), len(su["uses"])))] +
                   [["Total sources", su["total_sources"], "Total uses", su["total_uses"]]], [52, 15, 52, 15], {2: X.NUM, 4: X.NUM})
        diff_ = su["total_sources"] - su["total_uses"]
        cm.cell(row=rr, column=1, value=("Excluded (read for one year only): " + ", ".join(su["excluded"]) + ". " if su.get("excluded") else "")
                + (f"Sources - uses = {diff_:,.2f} Cr: does NOT reconcile - check the spread." if abs(diff_) > max(0.05, 0.005 * su["total_sources"])
                   else "Sources and uses reconcile.")).font = X.F_NOTE
        rr += 1
    for wsx in [wb[n_] for n_ in ("Projections & DSCR", "Repayment Schedule", "Scenarios", "WC Assessment & DP", "Security & Pricing",
                                  "RBI EWS & Regulatory", "CMA Analysis") if n_ in wb.sheetnames]:
        wsx.sheet_view.showGridLines = False
        for row in wsx.iter_rows(min_row=3):
            for cell in row:
                if cell.alignment is None or not cell.alignment.wrap_text:
                    cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str) and len(cell.value) > 40)
