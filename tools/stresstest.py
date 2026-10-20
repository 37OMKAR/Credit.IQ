"""Stress test: blank or corrupt every input line in turn and rebuild every output.

    python tools/stresstest.py none   # each line left unread in the latest year: all 8 reports + workbook, look for crashes and new '0.00' artefacts
    python tools/stresstest.py vals   # each line read as 0, negative or x100: prints the case status (a misread should not pass silently)
Slow (several minutes); not part of the quick self-test."""
import sys, copy, io, re, json, traceback, os, tempfile
HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, HERE)
os.environ["CREDITINTEL_DB"] = os.path.join(tempfile.mkdtemp(), "stress.db")
from core import pdf_reader as PR, extractor as EX, pipeline as PL, checker as CK, report_types as RT, report_excel as RX, schema as S
import docx
pages = PR.read_pdf(open(HERE + "/samples/Sample_Annual_Report_DeccanPolyfab_FY25.pdf", "rb").read())
loc = PR.locate_sections(pages)
raw = EX.heuristic_extract(pages, loc)
notes = EX.heuristic_notes_review(pages, loc)
Y = raw["years"]
ZERO = re.compile(r"[^\n|]{0,70}(₹ ?-?0\.0+ Cr|INR 0\.0+|\b0\.0+ Cr|\(0\.0%\)|\b0\.00x|\b0\.0%|\b0 days|\b0 d\b)[^\n|]{0,40}")
CAM = [k for k in RT.REPORT_TYPES if k.startswith("Credit Appraisal")][0]

def text_of(b):
    d = docx.Document(io.BytesIO(b)); t = [p.text for p in d.paragraphs]
    for tb in d.tables:
        for row in tb.rows: t.append(" | ".join(c.text for c in row.cells))
    return "\n".join(t)

def run(data, full):
    b = PL.analyse(data, Y, "Fuzz Co", notes=copy.deepcopy(notes), meta={"derived_fields": raw.get("derived")})
    n = PL.write_narrative(b)
    n, b["checks"] = CK.run(b, n)
    out = {"b": b, "n": n}
    if full:
        out["txt"] = text_of(RT.build(CAM, copy.deepcopy(b), n, "T", "HO"))
        RX.build(b)
        for name in RT.REPORT_TYPES:
            if name != CAM: RT.build(name, copy.deepcopy(b), n, "T", "HO")
    return out

base = run(copy.deepcopy(raw["data"]), True)
base_z = {m.group(0) for m in ZERO.finditer(base["txt"])}
print("baseline zero-like:", len(base_z), "status", base["b"]["integrity"]["status"]); sys.stdout.flush()
mode = sys.argv[1] if len(sys.argv) > 1 else "none"
keys = [k for k, *_ in S.ALL_ITEMS]
for k in keys:
    muts = {"none": [("None-latest", {Y[0]: None})],
            "vals": [("zero", {Y[0]: 0.0}), ("neg", {Y[0]: -abs((raw['data'][Y[0]].get(k) or 10.0))}), ("x100", {Y[0]: (raw['data'][Y[0]].get(k) or 1.0) * 100})]}[mode]
    for tag, ch in muts:
        d = copy.deepcopy(raw["data"])
        for y, v in ch.items(): d[y][k] = v
        try:
            r = run(d, mode == "none")
        except Exception as e:
            print(f"CRASH {k} {tag}: {type(e).__name__}: {e}"); traceback.print_exc(limit=-3); sys.stdout.flush(); continue
        if mode == "none":
            new = {m.group(0) for m in ZERO.finditer(r["txt"])} - base_z
            for z in sorted(new)[:6]:
                print(f"ZERO {k} {tag}: {z.strip()}")
        ig = r["b"]["integrity"]
        print(f"ok {k} {tag} status={ig['status']} grade={str(r['b']['forensics']['grade'])[:5]} rec={r['n'].get('recommendation','')[:20]}")
        sys.stdout.flush()
