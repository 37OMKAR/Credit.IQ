"""Self-test: proves the logic checker and the web-source checker catch what they must.

Run from the project folder:   python tools/selftest.py
Uses only the fictional sample annual report and a tiny local test website (no internet, no real borrower data).

It checks that
  1. the sample analysis passes every logic check and all 8 report types build;
  2. deliberately broken figures / ratio engine / grading are caught (the checker must FAIL them);
  3. an invented figure in AI text is removed while real figures stay;
  4. a lenient AI recommendation is overridden by the policy rule;
  5. web results are rejected when the link was never returned by the search engine, the page is about another
     company, the headline's claim is not on the page, a rating grade differs, or a number is not next to its label.
"""
import copy
import io
import os
import subprocess
import sys
import tempfile
import time

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, HERE)

from core import checker as CK  # noqa: E402
from core import extractor as EX  # noqa: E402
from core import pdf_reader as PR  # noqa: E402
from core import pipeline as PL  # noqa: E402
from core import ratios as R  # noqa: E402
from core import report_types as RT  # noqa: E402
from core import verify as V  # noqa: E402

FAILS = []


def ok(cond, msg):
    print(("  PASS  " if cond else "  FAIL  ") + msg)
    if not cond:
        FAILS.append(msg)


def results(checks):
    return {c["check"]: c["result"] for c in checks}


def sample_bundle():
    pdf = os.path.join(HERE, "samples", "Sample_Annual_Report_DeccanPolyfab_FY25.pdf")
    pages = PR.read_pdf(open(pdf, "rb").read())
    loc = PR.locate_sections(pages)
    r = EX.heuristic_extract(pages, loc)
    return PL.analyse(r["data"], r["years"], "Deccan Polyfab Industries Ltd", notes=EX.heuristic_notes_review(pages, loc)), r


print("1. Sample analysis")
b, raw = sample_bundle()
narr = PL.write_narrative(b)
_, chk = CK.run(b, narr)
fails = [c for c in chk if c["result"] == "FAIL"]
ok(not fails, f"sample passes all logic checks ({len(chk)} checks)" + (f" - failed: {[c['check'] for c in fails]}" if fails else ""))
for name in RT.REPORT_TYPES:
    out = RT.build(name, copy.deepcopy(b), narr, "Self test", "HO")
    ok(len(out) > 5000, f"report builds: {name}")

print("2. Deliberate errors must be caught")
d2 = copy.deepcopy(raw["data"])
y0 = raw["years"][0]
d2[y0]["total_assets"] = (d2[y0]["total_assets"] or 0) * 1.3
b2 = PL.analyse(d2, raw["years"], "Broken Co")
ok(results(CK.run(b2, PL.write_narrative(b2))[1]).get(f"Balance sheet tallies - {y0}") == "FAIL", "balance sheet that does not tally -> FAIL")

d3 = copy.deepcopy(raw["data"])
d3[y0]["revenue"] = None
b3 = PL.analyse(d3, raw["years"], "Missing Co")
ok(results(CK.run(b3, PL.write_narrative(b3))[1]).get("Core statement lines read") == "FAIL" and b3["forensics"]["grade"].startswith("NR"),
   "missing revenue -> FAIL and not rated")

b4 = copy.deepcopy(b)
b4["ratios"][y0]["current_ratio"] = (b4["ratios"][y0]["current_ratio"] or 1) * 1.5  # simulate a bug in the ratio engine
r4 = results(CK.run(b4, narr)[1])
ok("FAIL" in (r4.get("Key ratios recomputed independently agree with the ratio engine"),
              r4.get("TOL/TNW, ROCE, asset turnover, interest cover and current ratio recomputed for every year")), "wrong ratio from the engine -> FAIL")

b5 = copy.deepcopy(b)
b5["forensics"]["grade"] = "CI-1 - Very low risk"  # simulate a grading bug
ok(results(CK.run(b5, narr)[1]).get("Grade matches the score band") == "FAIL", "grade that does not match the score -> FAIL")

b6 = copy.deepcopy(b)
b6["ratios"][y0]["icr"] = 999.0
b6["data"][y0]["finance_cost"] = 0.0
ok(results(CK.run(b6, narr)[1]).get("Not-meaningful ratios are suppressed (shown as n.a. / n.m.)") == "FAIL",
   "interest cover shown on nil interest -> FAIL")

print("3. AI text")
rev = b["data"][y0]["revenue"]
ai = {"_source": "AI", "recommendation": narr["recommendation"],
      "executive_summary": f"Revenue was INR {rev:,.2f} Crore. The order book stands at INR 4,321.9 Crore. Margins are under pressure."}
clean, chk7 = CK.run(b, ai)
ok("4,321.9" not in clean["executive_summary"] and f"{rev:,.2f}" in clean["executive_summary"], "invented figure removed, real figure kept")

print("4. Policy rule beats a lenient AI recommendation")
clean8, _ = CK.run(b, dict(narr, recommendation="Approve"))
ok(clean8["recommendation"] == "Refer / further diligence" if (b["forensics"]["score"] < 55) else True, "AI 'Approve' overridden when the rule requires Refer")

print("5. Web sources (local test site)")
site = tempfile.mkdtemp()
open(os.path.join(site, "real.html"), "w").write("<h1>Deccan Polyfab Industries wins export order from Kenya</h1><p>Deccan Polyfab Industries Ltd said ... "
                                                  "Rated CARE A-/Stable. Scope 1 emissions 2,345.6 tCO2e. Revenue from operations 1,350.00 crore.</p>")
open(os.path.join(site, "other.html"), "w").write("<h1>Deccan Plastics Ltd results</h1><p>Deccan Plastics Ltd posts revenue of 2,345.6 crore</p>")
port = 8771
srv = subprocess.Popen([sys.executable, "-m", "http.server", str(port)], cwd=site, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
time.sleep(1.2)
# links come from search results / AI output: internal addresses must never be fetched (SSRF)
ok(all(V.url_blocked(u) for u in (f"http://localhost:{port}/real.html", "http://127.0.0.1/", "http://169.254.169.254/latest/meta-data/",
                                   "http://10.0.0.5/", "http://[::1]/", "file:///etc/passwd", "http://user:pw@example.com/")),
   "internal / loopback / cloud-metadata / non-http links are refused")
import http.server  # noqa: E402
import threading  # noqa: E402


class _Redir(http.server.BaseHTTPRequestHandler):
    def do_GET(self):  # a 'public' page that redirects into the internal network
        self.send_response(302)
        self.send_header("Location", "http://127.0.0.1/admin")
        self.end_headers()

    def log_message(self, *a):
        pass


rs = http.server.HTTPServer(("127.0.0.1", 8772), _Redir)
threading.Thread(target=rs.serve_forever, daemon=True).start()
_ub = V.url_blocked
V.url_blocked = lambda u: "" if ":8772/" in u else _ub(u)  # pretend the first host is public
txt_, why_ = V.fetch_text("http://127.0.0.1:8772/page")
V.url_blocked = _ub
rs.shutdown()
ok(txt_ is None and "non-public" in why_, "redirect from a public page to an internal address is refused")
V.ALLOW_PRIVATE = True  # the rest of this section uses the local test site
try:
    B = f"http://localhost:{port}/"
    co = "Deccan Polyfab Industries Ltd"
    cites = [{"url": B + "real.html"}, {"url": B + "other.html"}]
    good, bad = V.verify_news([
        {"title": "Deccan Polyfab Industries wins export order from Kenya", "url": B + "real.html"},
        {"title": "Deccan Polyfab Industries raided by Enforcement Directorate", "url": B + "real.html"},
        {"title": "Deccan Polyfab results", "url": B + "other.html"},
        {"title": "Deccan Polyfab fraud", "url": B + "invented.html"}], co, None, cites)
    ok(len(good) == 1 and len(bad) == 3, "news: real story kept; unsupported claim, other company and invented link rejected")
    ok(V.verify_rating({"agency": "CARE", "rating": "A-", "url": B + "real.html"}, co, None, cites)[0] is not None, "rating written on the page -> kept")
    ok(V.verify_rating({"agency": "CARE", "rating": "AA", "url": B + "real.html"}, co, None, cites)[0] is None, "different rating grade -> rejected")
    w = V.verify_web_pcaf({"scope1_tco2e": 2345.6, "revenue_inr_cr": 2345.6, "scope3_tco2e": 777.7, "sources": [{"url": B + "real.html"}],
                           "_citations": cites}, co, None)
    ok(w["scope1_tco2e"] == 2345.6, "emission figure next to 'Scope 1' on the company's page -> kept")
    ok(w["revenue_inr_cr"] is None, "same number but not next to 'revenue' -> rejected")
    ok(w["scope3_tco2e"] is None, "number not on any page -> rejected")
finally:
    srv.terminate()
    V.ALLOW_PRIVATE = False

print("6. Second, independent reading (simulated readers - no internet, no cost)")
import json  # noqa: E402
from core import dualread as D  # noqa: E402
pdf = open(os.path.join(HERE, "samples", "Sample_Annual_Report_DeccanPolyfab_FY25.pdf"), "rb").read()
pages = PR.read_pdf(pdf)
loc = PR.locate_sections(pages)


def reading(scale=1.0, unit="INR Crore", tweak=None):
    data = {y: {k: (None if raw["data"][y].get(k) is None else raw["data"][y][k] / scale) for k in D.COMPARE_KEYS} for y in raw["years"]}
    if tweak:
        tweak(data)
    return {"unit_in_document": unit, "years": raw["years"], "data": data}


def simulate(read2, read3_values=None, ext=None):
    def fake_post(api_key, body, case_id, task, timeout=180):
        return json.dumps(read2) if task == "read2_vision" else json.dumps({"unit_in_document": read2["unit_in_document"], "values": read3_values or {}})
    D._post = fake_post
    return D.run("simulated", pdf, loc, copy.deepcopy(ext or raw), PR.detect_unit)


rev = raw["data"][y0]["revenue"]
r = simulate(reading(0.01, "Rs in Lakhs"))
ok(r["summary"]["confirmed"] == r["summary"]["lines"], "both readers agree (second one in lakhs) -> every line confirmed")
r = simulate(reading(tweak=lambda d: d[y0].__setitem__("revenue", rev + 100)), {f"revenue|{y0}": rev})
row = next(x for x in r["rows"] if x["key"] == "revenue" and x["year"] == y0)
ok(row["status"] == "Confirmed by tie-breaker" and row["final"] == rev, "second reader misreads revenue -> tie-breaker keeps the right figure")
wrong = copy.deepcopy(raw)
wrong["data"][y0]["pat"] = raw["data"][y0]["pat"] * 10
r = simulate(reading(), {f"pat|{y0}": raw["data"][y0]["pat"]}, ext=wrong)
D.apply_corrections(wrong, r)
ok(wrong["data"][y0]["pat"] == raw["data"][y0]["pat"], "first reading wrong -> corrected when two readers agree")
r = simulate(reading(tweak=lambda d: d[y0].__setitem__("revenue", rev + 100)), {f"revenue|{y0}": rev + 55})
ok(r["summary"]["disputed"] == 1, "tie-breaker cannot settle -> left as disputed for the officer")
bd = copy.deepcopy(b)
bd.setdefault("meta", {})["dual"] = r
ok(results(CK.run(bd, narr)[1]).get("Figures confirmed by a second, independent reading") == "FAIL", "disputed revenue -> checker FAIL (report DRAFT)")
r = simulate(reading(0.01, "INR Crore"))
ok(r["summary"].get("unit_mismatch") == 0.01, "readers disagree on the unit -> unit mismatch reported")
ok(PR.detect_unit("(Amount in Rupees)") == 1e-7 and PR.detect_unit("Rs. in Lakhs") == 0.01 and PR.detect_unit("₹ in hundreds") == 1e-5,
   "units: rupees, lakhs and hundreds recognised")
bu = copy.deepcopy(b)
bu["data"][y0]["total_assets"] = bu["data"][y0]["total_assets"] * 1e7
ok(results(CK.run(bu, narr)[1]).get("Size plausible for the unit read (rupees / thousands / lakhs / crore)") == "FAIL",
   "figures in rupees read as crore -> FAIL")

print("7. The external audit's regression cases (SECI-style consolidated test document)")
sys.path.insert(0, os.path.join(HERE, "tools"))
import make_consolidated_sample as MC  # noqa: E402
from core import credit_tools as CT  # noqa: E402
cpath = os.path.join(tempfile.mkdtemp(), "consolidated_sample.pdf")
MC.build(cpath)
cpages = PR.read_pdf(open(cpath, "rb").read())
cloc = PR.locate_sections(cpages, prefer="standalone")
cext = EX.heuristic_extract(cpages, cloc)
cy = cext["years"][0]


def near(a, b_):
    return a is not None and abs(a - b_) <= max(0.02, 0.001 * abs(b_))


ok(all(near(cext["data"][y_].get(k), v) for y_, exp in MC.EXPECTED.items() for k, v in exp.items() if k in ("revenue", "total_current_liabilities")),
   "revenue and current liabilities extracted for both years")
ok(near(cext["data"][cy].get("short_term_borrowings"), 50.75), "current borrowings on the second balance-sheet page captured")
ok(near(cext["data"][cy].get("finance_cost"), 10.934), "finance cost taken from the P&L (10.93 Cr), not cash-flow 'interest on loan' (3.02 Cr)")
ok(cloc["basis"] == "consolidated" and "consolidated statements only" in cloc["basis_note"], "consolidated-only document labelled consolidated, not standalone")
ok(near(cext["data"][cy].get("capex"), 690.05), "capex includes PPE, CWIP investment and capital advances")
cb = PL.analyse(cext["data"], cext["years"], MC.CO, notes=EX.heuristic_notes_review(cpages, cloc), statement_type=cloc["basis"],
                screening={"attest": {}})
ok(near(cb["data"][cy]["ebitda"], 519.16), "EBITDA excludes other income and the JV share (519.16 Cr)")
ok(near(cb["data"][cy]["total_debt"], 300.80), "total debt includes current borrowings and leases (300.80 Cr)")
ok(near(cb["data"][cy]["nwc"], 471.28) and near(cb["ratios"][cy]["current_ratio"], 1.10), "NWC 471.28 Cr and current ratio 1.10x")
miss = copy.deepcopy(cext["data"])
from core import schema as SC  # noqa: E402
for k_, _l, sec_, _s in SC.ALL_ITEMS:  # current liabilities genuinely unknown: the total AND every line it could be rebuilt from
    if sec_ == "Current Liabilities":
        miss[cy][k_] = None
mb = PL.analyse(miss, cext["years"], MC.CO)
ok(all(mb["ratios"][cy].get(k) is None for k in ("current_ratio", "quick_ratio", "nwc_cr", "mpbf2")) and mb["data"][cy]["nwc"] is None,
   "a blank mandatory input -> dependent ratios withheld, never computed as zero")
neg = copy.deepcopy(cext["data"])
neg[cy]["total_current_liabilities"] = 6000.0
nb = PL.analyse(neg, cext["years"], MC.CO)
ok(nb["data"][cy]["nwc"] < 0 and nb["ratios"][cy]["current_ratio"] < 1, "current liabilities above current assets -> negative NWC reported")
pz = dict(CT.proposal(cb), cc_amount=0.0)
cbz = copy.deepcopy(cb)
cbz["data"][cy]["short_term_borrowings"] = 0.0
cbz["data"][cy]["current_maturities_ltd"] = 0.0
ok(CT.drawing_power(cbz, pz)["available"] == 0.0, "zero working-capital limit -> available drawing power is zero")
nr, _ = CK.run(mb, dict(PL.write_narrative(mb), recommendation="Approve"))
ok(not nr["recommendation"].startswith("Approve"), "a failed check -> no approval recommendation")
cn = EX.heuristic_notes_review(cpages, cloc)
ok(str(cn["auditor_opinion"]).startswith("not verified"), "auditor's report not supplied -> audit opinion 'not verified'")
ca, _ = CK.run(cb, dict(PL.write_narrative(cb), recommendation="Approve"))
ok(ca["recommendation"] == "Refer / further diligence", "unverified opinion / old statements -> approval withheld even when the figures are good")
pd_ = dict(CT.proposal(b), tl_amount=200.0, tl_rate=0.1, tl_tenor_years=5)
sc_ = CT.scenarios(b, pd_)
ok(len({round(x["avg_dscr"] or 0, 4) for x in sc_}) == len(sc_), "scenarios respond to their growth / margin / rate assumptions")

print("8. Required inputs - nothing silently read as zero (audit P0)")
y0s, y1s = b["years"][0], b["years"][1]
d8 = copy.deepcopy(raw["data"])
d8[y0s]["capex"] = None
b8 = PL.analyse(d8, raw["years"], "T")
ok(b8["data"][y0s].get("fcf") is None and b8["ratios"][y0s].get("fcf_cr") is None, "capex not read -> free cash flow withheld (not CFO - 0)")
d8 = copy.deepcopy(raw["data"])
d8[y1s]["revenue"] = None
b8 = PL.analyse(d8, raw["years"], "T")
_r8, a8 = CT.projections(b8, dict(CT.proposal(b8), growth=None))
ok(a8.get("status") == "NOT COMPUTED" and not _r8, "no historical growth and none entered -> projections NOT COMPUTED (no default growth)")
ok(EX._dash_value("revenue") is None and EX._dash_value("goodwill") == 0, "printed dash: blank for a required line, nil only for an optional line")
d8 = copy.deepcopy(raw["data"])
d8[y1s]["trade_receivables"] = None
d8[y0s]["finance_cost"] = None
b8 = PL.analyse(d8, raw["years"], "T")
na8 = {t["test"] for t in b8["forensics"]["not_assessed"]}
ok({"Receivables growing much faster than revenue", "Implied interest rate unusually low"} <= na8
   and not any(f["title"] in na8 for f in b8["forensics"]["flags"]), "forensic tests without their inputs -> NOT ASSESSED, and no flag raised on a zero")
ok(b8["forensics"]["models"]["Beneish M-score"]["status"] == "NOT COMPUTED" and b8["forensics"]["beneish"] is None,
   "Beneish with a missing core input -> NOT COMPUTED (no score)")
d8 = copy.deepcopy(raw["data"])
d8[y0s]["trade_receivables"] = None
d8[y0s]["ppe"] = None
b8 = PL.analyse(d8, raw["years"], "T")
ok(CT.drawing_power(b8, CT.proposal(b8))["dp"] is None and CT.security_cover(b8, CT.proposal(b8))["cover"] is None,
   "receivables / fixed assets not read -> drawing power and security cover NOT COMPUTED")
CK.run(mb, PL.write_narrative(mb))
ok(mb["integrity"]["status"] == "BLOCKED" and not mb["integrity"]["final_recommendation_allowed"] and mb["integrity"]["missing_required_fields"],
   "required line missing -> case status BLOCKED, no recommendation allowed")
CK.run(b, narr)
ok(b["integrity"]["status"] in ("READY", "READY_WITH_WARNINGS") and b["integrity"]["final_recommendation_allowed"],
   "complete sample -> case status not blocked")
for name in RT.REPORT_TYPES:
    out = RT.build(name, copy.deepcopy(mb), PL.write_narrative(mb), "Self test", "HO")
ok(True, "all report types build for a BLOCKED case with missing inputs")

print("9. Access control and audit trail (temporary database)")
from core import db as DB  # noqa: E402
DB.DB_PATH = os.path.join(tempfile.mkdtemp(), "t.db")
DB.init()
DB.create_user("officer.a", "Correct#2026")
DB.create_user("officer.b", "Correct#2026")
for _ in range(DB.MAX_FAILS):
    DB.verify("officer.a", "wrong-password")
ok(DB.login_locked("officer.a") > 0 and DB.verify("officer.a", "Correct#2026") is None,
   f"{DB.MAX_FAILS} wrong passwords -> username locked, even the right password is refused for {DB.LOCK_MINUTES} min")
ok(DB.verify("officer.b", "Correct#2026") is not None, "another user is not affected by the lock")
DB.set_password("officer.a", "Correct#2026")
ok(DB.verify("officer.a", "Correct#2026") is not None, "an administrator's password reset clears the lock")
ua, ub, adm = {"username": "officer.a", "role": "officer"}, {"username": "officer.b", "role": "officer"}, {"username": "x", "role": "admin"}
cases_dir = os.path.join(HERE, "cases")
ok(DB.can_access(ua, "officer.a") and not DB.can_access(ub, "officer.a") and DB.can_access(adm, "officer.a"),
   "an officer opens only their own cases; an administrator opens all")
ok(not DB.can_access(ua, "officer.a", os.path.join(HERE, "core", "db.py"), cases_dir)
   and DB.can_access(ua, "officer.a", os.path.join(cases_dir, "officer.a", "c.json"), cases_dir),
   "a stored file path outside the application's own folder is refused")
CASE_A = "officer.a-20261020101010000000"  # case ids carry their owner (as the app creates them)
DB.add_figure_changes(CASE_A, "T", "officer.a", [{"year": "FY2025", "key": "revenue", "line": "Revenue", "old": 135.0, "new": 1350.0}],
                      "per audited P&L p.42")
fc = DB.figure_changes(CASE_A)
ok(len(fc) == 1 and fc[0]["old_value"] == 135.0 and fc[0]["new_value"] == 1350.0 and fc[0]["reason"] and fc[0]["username"] == "officer.a",
   "a figure change is stored with who, when, old, new and reason")

print("10. Hidden-mistake review: derived totals, sub-totals, Tandon, growth")
from core import schema as SCH  # noqa: E402
lg = {}
yy = SCH.fill_derived({"pbt": 100.0, "pat": None, "tax_expense": None, "share_capital": 10.0, "other_equity": None}, lg)
ok(yy.get("pat") is None and yy.get("total_equity") is None,
   "tax not read -> PAT is NOT set to PBT; other equity not read -> net worth is NOT set to share capital")
lg = {}
yy = SCH.fill_derived({"pbt": 100.0, "pat": 90.0, "tax_expense": 30.0, "total_assets": 500.0}, lg)
ok(yy["tax_expense"] == 30.0 and "total_equity_and_liabilities" in lg and lg["total_equity_and_liabilities"].startswith("copied"),
   "a tax figure that was read is never overwritten; a copied total is recorded as derived")
d10 = copy.deepcopy(raw["data"])
d10[y0s]["tax_expense"], d10[y0s]["total_equity_and_liabilities"] = None, None
b10 = PL.analyse(d10, raw["years"], "T")
r10 = results(CK.run(b10, PL.write_narrative(b10))[1])
ok(r10.get(f"PBT - tax = PAT - {y0s}") == "WARN" and r10.get(f"Balance sheet tallies - {y0s}") == "WARN",
   "tie-outs on a derived figure are reported 'cannot be verified' - never 'passed'")
for k10, f10 in (("trade_receivables", 100), ("revenue", 100), ("other_expenses", 10)):
    d10 = copy.deepcopy(raw["data"])
    d10[y0s][k10] = d10[y0s][k10] * f10
    b10 = PL.analyse(d10, raw["years"], "T")
    CK.run(b10, PL.write_narrative(b10))
    ok(b10["integrity"]["status"] == "BLOCKED", f"one line misread x{f10} ({k10}) -> its lines no longer add to the printed total -> BLOCKED")
ok(abs(b["ratios"][y0s]["etr"] - (b["data"][y0s]["pbt"] - b["data"][y0s]["pat"]) / b["data"][y0s]["pbt"]) < 1e-9,
   "effective tax rate is (PBT - PAT) / PBT - independent of how the tax line was read")
d10 = copy.deepcopy(raw["data"])
d10[y0s]["current_maturities_ltd"] = 20.0
b10 = PL.analyse(d10, raw["years"], "T")
c10 = b10["data"][y0s]
exp10 = 0.75 * c10["total_current_assets"] - (c10["total_current_liabilities"] - (c10["short_term_borrowings"] - 20.0))
ok(abs(b10["ratios"][y0s]["mpbf2"] - exp10) < 1e-6 and abs(CT.wc_methods(b10, CT.proposal(b10))["tandon2"] - exp10) < 1e-6,
   "Tandon II: term-loan instalments due within a year stay in other current liabilities (not treated as bank finance)")
d10 = copy.deepcopy(raw["data"])
d10[y1s]["pat"], d10[y1s]["tax_expense"] = -10.0, 0.0
b10 = PL.analyse(d10, raw["years"], "T")
ok(b10["ratios"][y0s]["pat_growth"] is None, "growth on a loss-making prior year is n.m. (not '-1,210%')")
d10 = copy.deepcopy(raw["data"])
d10[y0s]["long_term_borrowings"] = None
b10 = PL.analyse(d10, raw["years"], "T")
ok(results(CK.run(b10, PL.write_narrative(b10))[1]).get("Borrowings read for the latest year") == "FAIL",
   "borrowings read last year but missing this year -> FAIL (never a silent 'debt-free' D/E of 0.00x)")
from core import watchlist as WLm  # noqa: E402
h10, _ = WLm.screen([{"name": "Ramesh Kumar Sharma", "role": "Director", "id": ""}],
                    [{"list_name": "Wilful defaulter", "name": "Ramesh Kumar Sharma", "identifier": "", "directors": ""}])
ok(h10 and h10[0]["severity"] == "Medium" and "name only" in h10[0]["match"],
   "an individual matched on name alone is 'verify identity' (Medium), not a High hit")

print("11. External audit (2): unknown vs zero, drawing power, PCAF, nil policy, partial web pages, audit trail")
pages_s = PR.read_pdf(open(os.path.join(HERE, "samples", "Sample_Annual_Report_DeccanPolyfab_FY25.pdf"), "rb").read())
loc_s = PR.locate_sections(pages_s)
# 1. unknown versus zero - every reading keeps its column and its source token
lv = EX._line_values("Other current assets 0 12.50", True, 2, with_tokens=True)
ok(lv == [(12.5, "12.50", "value")], "decimal page: a note number before the first value is not a value")
lv = EX._line_values("Other current assets 12.50 0", True, 2, with_tokens=True)
ok(lv == [(12.5, "12.50", "value"), (0.0, "0", "zero")], "decimal page: a printed '0' is an explicit zero and keeps its column")
lv = EX._line_values("Trade payables 12 1,2A4.00 980.00", True, 2, with_tokens=True)
ok(lv == [(None, "1,2A4.00", "unreadable"), (980.0, "980.00", "value")],
   "an OCR-garbled figure is left blank in ITS column - last year's 980 never slides into this year")
lv = EX._line_values("Trade payables 1,O50.00 980.00", False, 2, with_tokens=True)
ok(lv[0] == (1050.0, "1,O50.00", "ocr-repaired") and lv[1][0] == 980.0, "an unambiguous OCR slip (O for 0) is repaired and marked for checking")
lv = EX._line_values("Borrowings - 120.00", True, 2, with_tokens=True)
ok(lv[0][0] == EX.DASH and lv[0][2] == "dash" and lv[1][0] == 120.0, "a printed dash is recorded as a dash, not as a number")
pv = raw.get("provenance") or {}
ok(pv and all(e.get("token") and e.get("kind") and e.get("page") for y_ in pv for lst in pv[y_].values() for e in lst),
   "every figure read by the rule-based reader keeps its printed token, how it was read and its page")


class _FakeLLM:
    provider = "test"

    def __init__(self, out):
        self.out = out

    def chat_json(self, *a, **k):
        return self.out


fake = _FakeLLM({"years": ["FY2025"], "unit_in_document": "INR Crore", "data": {"FY2025": {
    "revenue": 100, "trade_receivables": "-", "goodwill": "-", "contingent_liabilities": "N/A", "intangibles": "nil",
    "inventories": 0, "other_income": None}}})
lx = EX.llm_extract_financials(fake, pages_s, loc_s)["data"]["FY2025"]
lp = EX.llm_extract_financials(fake, pages_s, loc_s)["provenance"]["FY2025"]
ok(lx["trade_receivables"] is None and lp["trade_receivables"][0]["kind"].startswith("UNRESOLVED"),
   "AI reading: a dash on a required line stays blank and is marked UNRESOLVED")
ok(lx["goodwill"] == 0.0 and lx["contingent_liabilities"] is None and lx["intangibles"] == 0.0 and lx["inventories"] == 0.0
   and lx["other_income"] is None and lp["contingent_liabilities"][0]["kind"].startswith("not applicable"),
   "AI reading: dash on an optional line = nil, 'N/A' = blank, 'nil' and 0 = stated zero, missing = blank - each recorded")
# 2. drawing power: an entered 0 CC limit is respected
bdp = copy.deepcopy(b)
bdp["data"][y0s]["short_term_borrowings"], bdp["data"][y0s]["current_maturities_ltd"] = 80.0, 10.0
dp0 = CT.drawing_power(bdp, dict(CT.proposal(bdp), cc_amount=0.0))
dpn = CT.drawing_power(bdp, dict(CT.proposal(bdp), cc_amount=None))
ok(dp0["limit"] == 0.0 and dp0["available"] == 0.0, "CC limit entered as 0 with existing borrowings of 70 Cr -> limit 0, no drawing power")
ok(abs(dpn["limit"] - 70.0) < 1e-9 and dpn["available"] == min(dpn["dp"], 70.0), "CC limit not specified -> existing WC borrowings (70 Cr) used")
ok(CT.wc_methods(bdp, dict(CT.proposal(bdp), cc_amount=0.0))["requested"] == 0.0, "WC assessment: an entered 0 is the request, not 'existing'")
ok(CT.proposal(bdp)["cc_amount"] is None, "default CC limit is 'not specified', not 0")
# 3. PCAF: unreported emissions stay unknown
from core import pcaf as PCm  # noqa: E402
pu = PCm.calculate(10.0, "Business loans / unlisted equity", 400.0, 600.0)
ok(pu["financed_scope12"] is None and pu["financed_scope3"] is None and pu["data_quality_score"] is None,
   "no emissions reported and no sector factor -> financed emissions NOT computed (not 0)")
pz = PCm.calculate(10.0, "Business loans / unlisted equity", 400.0, 600.0, scope1=0.0, scope2=0.0)
ok(pz["financed_scope12"] == 0.0 and pz["data_quality_score"] == 2, "Scope 1 and 2 reported as zero -> financed 0, data quality 2")
p1 = PCm.calculate(10.0, "Business loans / unlisted equity", 400.0, 600.0, scope1=500.0)
ok(p1["financed_scope12"] is None and "Only Scope 1" in p1["warning"], "only Scope 1 reported -> not used as the reported Scope 1+2 total")
ok(PCm.calculate(10.0, "Business loans / unlisted equity", None, 600.0, scope1=1.0, scope2=1.0)["attribution_factor"] is None
   and PCm.calculate(None, "Business loans / unlisted equity", 400.0, 600.0)["attribution_factor"] is None,
   "equity not read or exposure not entered -> attribution factor not computed (not 0)")
from core import report_excel as RXm  # noqa: E402
import openpyxl as _ox  # noqa: E402
bpc = copy.deepcopy(b)
wbp = _ox.load_workbook(io.BytesIO(RXm.build(bpc, pcaf_inputs={"outstanding_cr": 10.0, "asset_class": "Business loans / unlisted equity"})))
w3 = wbp["PCAF Financed Emissions"]
ok(w3["B13"].value is None and w3["B14"].value is None and w3["B15"].value is None and "ISNUMBER(B13)" in str(w3["B21"].value)
   and "ISNUMBER(B15)" in str(w3["B23"].value) and "ISBLANK" in str(w3["B9"].value),
   "workbook: unreported emissions are blank inputs and every formula returns n.a. on a blank (no zero)")
# 4. nil policy
ok(set(R.NIL_POLICY) == R.OPTIONAL_NIL, "every line allowed to be nil has a documented nil policy")
dnil = copy.deepcopy(raw["data"])
dnil[y0s]["long_term_borrowings"], dnil[y0s]["short_term_borrowings"] = None, None
bnil = PL.analyse(dnil, raw["years"], "T")
ok(bnil["data"][y0s]["total_debt"] is None and bnil["ratios"][y0s]["de"] is None and bnil["ratios"][y0s]["debt_ebitda"] is None,
   "borrowings not read while finance cost is material -> total debt and leverage ratios withheld (not 'debt-free')")
ok(results(CK.run(bnil, PL.write_narrative(bnil))[1]).get("Lines treated as nil are credibly nil") == "WARN",
   "... and the officer is asked to confirm those lines")
dnil = copy.deepcopy(raw["data"])
for k_ in ("cost_of_materials", "purchases_stock", "change_in_inventories"):
    dnil[y0s][k_] = None
bnil = PL.analyse(dnil, raw["years"], "T")
ok(bnil["ratios"][y0s]["inventory_days"] is None and "cost_of_materials" in R.unconfirmed_nil(bnil["data"][y0s]),
   "inventory held but no cost-of-goods line read -> inventory days withheld (never 'service business')")
okn = copy.deepcopy(raw["data"])
bok = PL.analyse(okn, raw["years"], "T")
ok(not R.unconfirmed_nil(bok["data"][y0s]), "complete sample: no optional line wrongly blocked")
# 5. web pages: partial downloads are never evidence; every response is closed
closed = []


class _Resp:
    status_code, headers, encoding = 200, {"content-type": "text/html"}, "utf-8"

    def __init__(self, n):
        self.n = n

    def iter_content(self, size):
        for _ in range(self.n):
            yield b"<p>Deccan Polyfab Industries revenue 1,350.00</p>" * 40

    def close(self):
        closed.append(1)


_g0, _mb0 = V._get, V.MAX_BYTES
try:
    V._CACHE.clear()
    V.MAX_BYTES = 20000
    V._get = lambda u, t: _Resp(50)
    t_big, why_big = V.fetch_text("http://example.org/big")
    V._get = lambda u, t: _Resp(1)
    t_ok, _w = V.fetch_text("http://example.org/small")
    hdr = _Resp(1)
    hdr.headers = {"content-type": "text/html", "content-length": str(10 ** 9)}
    V._get = lambda u, t: hdr
    t_hdr, why_hdr = V.fetch_text("http://example.org/declared-big")
finally:
    V._get, V.MAX_BYTES = _g0, _mb0
    V._CACHE.clear()
ok(t_big is None and "not used as evidence" in why_big and t_hdr is None and "too large" in why_hdr and t_ok and "1,350.00" in t_ok,
   "oversized page (by size or declared length) -> rejected, not parsed from the part that was read; a normal page is read")
ok(len(closed) == 3, "every web response is closed (rejected or read)")
# 6. audit trail: ownership, identity, reason, timestamps, append-only, report = database
import sqlite3  # noqa: E402
CASE_B = "officer.b-20261020101010000000"
try:
    DB.add_figure_changes(CASE_A, "T", "officer.b", [{"year": "FY2025", "key": "pat", "line": "PAT", "old": 1.0, "new": 2.0}], "x")
    other_ok = False
except PermissionError:
    other_ok = True
ok(other_ok, "an officer cannot change figures on another officer's case")
try:
    DB.add_figure_changes("officer.a-x-20261020101010000000", "T", "officer.a", [{"year": "Y", "key": "k", "line": "l", "old": 1, "new": 2}], "x")
    pref_ok = False
except PermissionError:
    pref_ok = True
ok(pref_ok, "a case id that only starts with the user's name is not treated as theirs")
for args_, exc_ in (((CASE_A, "T", "officer.a", [{"year": "Y", "key": "k", "line": "l", "old": 1, "new": 2}], "  "), ValueError),
                    ((CASE_A, "T", "", [{"year": "Y", "key": "k", "line": "l", "old": 1, "new": 2}], "reason"), ValueError)):
    try:
        DB.add_figure_changes(*args_)
        good_ = False
    except exc_:
        good_ = True
    ok(good_, "a figure change without " + ("a reason" if args_[2] else "an identified user") + " is refused")
DB.create_user("boss", "Correct#2026", role="admin")
DB.add_figure_changes(CASE_B, "T", "boss", [{"year": "FY2025", "key": "pat", "line": "PAT", "old": 1.0, "new": 2.0}], "admin correction")
fc_b = DB.figure_changes(CASE_B)
ok(len(fc_b) == 1 and fc_b[0]["username"] == "boss", "an administrator's correction is recorded under the administrator's name")
import datetime as _dtm  # noqa: E402
ok(all(_dtm.datetime.fromisoformat(r_["ts"]) for r_ in DB.figure_changes(CASE_A)), "every change has a valid timestamp")
blocked_sql = 0
for sql_ in ("UPDATE figure_changes SET new_value=0", "DELETE FROM figure_changes"):
    try:
        with sqlite3.connect(DB.DB_PATH) as con_:
            con_.execute(sql_)
    except sqlite3.DatabaseError:
        blocked_sql += 1
ok(blocked_sql == 2 and len(DB.figure_changes(CASE_A)) == 1, "the audit trail cannot be edited or deleted, even directly in the database")
brec = copy.deepcopy(b)
dx_, _xx = PL.outputs(brec, PL.write_narrative(brec))
summ = PL.case_summary(brec)
cid = DB.add_case("officer.a", "T", summ, 0, 0, 0, "d", "x", "j", "s")
row_ = next(r_ for r_ in DB.cases("officer.a") if r_["id"] == cid)
import docx as _dx  # noqa: E402
cover = "\n".join(c_.text for t_ in _dx.Document(io.BytesIO(dx_)).tables for r__ in t_.rows for c_ in r__.cells)
ok(row_["recommendation"] == brec["final_narrative"]["recommendation"] and row_["recommendation"] in cover
   and row_["grade"] == brec["forensics"]["grade"] and abs(row_["score"] - brec["forensics"]["score"]) < 1e-9,
   "report-to-database: the stored recommendation, grade and score are exactly those printed in the memorandum")
ok(all(r_["username"] == "officer.a" for r_ in DB.cases("officer.a")), "an officer's case list holds only their own cases")

print("12. Unit and currency are read, never assumed (case: Pakistani statements in whole rupees read as crore)")
ok(PR.detect_unit_info("Rupees")[0] == 1e-7 and PR.detect_unit_info("PKR")[0] == 1e-7, "a bare 'Rupees' unit = whole rupees (not crore)")
ok(PR.detect_unit_info("Statement of profit and loss for the year")[0] is None, "no unit statement -> unit unknown (not silently crore)")
ok(PR.detect_currency("Karachi, Pakistan. Securities and Exchange Commission of Pakistan. Pakistan Stock Exchange. Amount in Rupees")[0] == "PKR",
   "Pakistani statements -> currency PKR detected")
ok(PR.detect_currency("Amounts in USD thousands")[0] == "USD" and
   PR.detect_currency("India. Companies Act 2013. Unhedged foreign currency exposure amount of USD 1.2 million")[0] is None,
   "a USD presentation unit is detected; a USD exposure note in an Indian report is not")
bu = PL.analyse(copy.deepcopy(raw["data"]), raw["years"], "T", meta={"unit_status": "NOT STATED"})
CK.run(bu, PL.write_narrative(bu))
ok(bu["integrity"]["status"] == "BLOCKED", "unit not stated in the document -> case BLOCKED until the officer sets it")
bc = PL.analyse(copy.deepcopy(raw["data"]), raw["years"], "T", meta={"currency": {"code": "PKR", "evidence": "test"}, "unit_status": "Rs. in crore"})
rc = results(CK.run(bc, PL.write_narrative(bc))[1])
ok(rc.get("Statements are in Indian rupees") == "FAIL" and bc["integrity"]["status"] == "BLOCKED", "statements in PKR -> FAIL, case BLOCKED")
fake_u = _FakeLLM({"years": ["FY2015"], "unit_in_document": "Rupees", "data": {"FY2015": {"revenue": 4094937171, "total_assets": 2835665454}}})
eu = EX.llm_extract_financials(fake_u, pages_s, loc_s)
ok(abs(eu["data"]["FY2015"]["revenue"] - 409.4937171) < 1e-3 and eu["unit_status"] == "Rupees",
   "AI reader says 'Rupees' -> Rs 4,094,937,171 becomes Rs 409.49 crore")

print()
if FAILS:
    print(f"SELF-TEST FAILED: {len(FAILS)} problem(s)")
    for f in FAILS:
        print("  -", f)
    sys.exit(1)
print("SELF-TEST PASSED - every check behaved as required.")
